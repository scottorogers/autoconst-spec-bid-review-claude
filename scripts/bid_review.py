#!/usr/bin/env python3
"""
spec-bid-review: read a spec / tender (one or more PDFs) and produce the estimator's
spec-review form as Excel. Every row cites Section + Page and carries a confidence flag.

Usage:
    py scripts/bid_review.py spec.pdf [more.pdf ...] [--region uk|us] [--trade all|electrical|mep|concrete|finishes] [--out outputs/spec-review.xlsx]
    (a pre-extracted pdftotext -layout .txt also works as input)

Regions:
    uk (default) - UK tenders: ITT / JCT / NEC, NBS work sections (Y61...), Uniclass, GBP, LADs, Employer / CA
    us           - US CSI MasterFormat specs and public-works bid documents (the original engine, unchanged)

Page numbers are PDF page numbers (1 = first sheet of the PDF), not the printed page number.

Tabs:
    Summary            - read this first: bid-killers with their values, top hidden costs, top RFIs
    Bid-Killers        - Div 00/01 commercial risks, one row each, key value pulled out
    Hidden Costs       - "no separate payment", "incidental", "at Contractor's expense"...
    Sole-Source        - named products classified sole-source / limited list / basis of design / or-equal
    Submittals         - register: type, what, timing
    Costly Items       - trade watchlists (config/watchlists.json), one row per item with every page
    RFI List           - price-relevant ambiguity, ranked, RFI question drafted
    Method & Notes

Nothing is invented: every row is a real sentence from the document. Missing bid-killers are
listed as NOT FOUND.
"""
import argparse, bisect, hashlib, json, os, re, shutil, subprocess, sys
from collections import Counter, OrderedDict
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
CONFIG = ROOT / "config" / "watchlists.json"

# =====================================================================================
# 1. Loading, header/footer removal, section tracking, sentences
# =====================================================================================

def find_pdftotext():
    """PDFTOTEXT env var wins (lets tests pin xpdf vs Poppler); otherwise whatever is on PATH."""
    exe = os.environ.get("PDFTOTEXT") or shutil.which("pdftotext")
    if not exe or not Path(exe).exists() and not shutil.which(exe):
        sys.exit("pdftotext was not found. Install Poppler (Windows: winget install oschwartz10612.Poppler; "
                 "macOS: brew install poppler; Linux: apt install poppler-utils) or set PDFTOTEXT to its full path.")
    return exe

def pdftotext_version(exe):
    try:
        r = subprocess.run([exe, "-v"], capture_output=True, text=True, timeout=20)
        line = (r.stderr or r.stdout).strip().splitlines()
        return line[0] if line else "unknown version"
    except (OSError, subprocess.SubprocessError):
        return "unknown version"

def pdf_to_text(pdf, workdir, exe):
    workdir.mkdir(parents=True, exist_ok=True)
    # suffix with a path hash so two inputs named the same (in different folders) don't collide;
    # hashlib, not hash(), so the name is the same on every run and work/ doesn't fill up
    out = workdir / f"{pdf.stem}-{hashlib.md5(str(pdf.resolve()).encode()).hexdigest()[:8]}.txt"
    r = subprocess.run([exe, "-layout", str(pdf), str(out)], capture_output=True, text=True)
    if r.returncode != 0 or not out.exists():
        msg = (r.stderr or "").strip().splitlines()
        sys.exit(f"pdftotext could not read {pdf.name} (exit {r.returncode}): {msg[-1] if msg else 'no details'}. "
                 "The PDF may be encrypted or damaged.")
    return out.read_text(encoding="utf-8", errors="replace")

# Different pdftotext builds (xpdf 4.x vs Poppler) emit different quote / ligature characters.
# Normalise them once so every pattern below only has to handle plain ASCII punctuation.
NORMALISE = str.maketrans({"‘": "'", "’": "'", "‚": "'", "‛": "'", "′": "'",
                           "“": '"', "”": '"', "„": '"', "″": '"',
                           " ": " ", " ": " ", " ": " ", "­": None,
                           "ﬁ": "fi", "ﬂ": "fl", "ﬀ": "ff", "ﬃ": "ffi", "ﬄ": "ffl"})

def normalise(text):
    return text.translate(NORMALISE)

FOOTER_SEC = re.compile(r"(?<!\d)(\d{2} \d{2} \d{2}(?:\.\d{2})?)\s*[-–—]\s*\d{1,3}\b|(?<!\d)(\d{5}(?:\.\d{2})?)\s*[-–—]\s*\d{1,3}\s+of\s+\d{1,3}", re.M)
# (kind, pattern) - kind decides how the label is written
HDR_PATTERNS = [
    ("csi", re.compile(r"^[ \t]*SECTION[ \t]+(\d{2}[ \t]?\d{2}[ \t]?\d{2}(?:\.\d{2})?)\b[ \t]*[-–—:]?[ \t]*([A-Z][A-Z0-9 ,&/()'\-]{2,60})?", re.M)),
    ("csi", re.compile(r"^[ \t]*SECTION[ \t]+(\d{5}(?:\.\d{2})?)\b[ \t]*[-–—:]?[ \t]*([A-Z][A-Z0-9 ,&/()'\-]{2,60})?", re.M)),
    ("roman", re.compile(r"^[ \t]*SECTION[ \t]+([IVXL]{1,6})\.[ \t]+([A-Z][A-Z0-9 ,&/()'\-]{2,60})", re.M)),
    ("division", re.compile(r"^[ \t]*DIVISION[ \t]+(\d{1,2})\b[ \t]*[-–—:]?[ \t]*([A-Z][A-Z0-9 ,&/()'\-]{2,60})?", re.M)),
    ("exhibit", re.compile(r"^[ \t]*EXHIBIT[ \t]+\"?([A-Z])\"?[ \t]*$", re.M)),
]
HDR_PATTERNS_US = HDR_PATTERNS
# UK: NBS / CAWS work sections ("SECTION Y61 LV cables", or a bare "Y61 LV CABLES AND WIRING" line)
# and Uniclass 2015 codes ("Ss_70_30_45 Low voltage distribution systems"). Labels kept verbatim.
HDR_PATTERNS_UK = HDR_PATTERNS_US + [
    ("verbatim", re.compile(r"^[ \t]*(?:WORK[ \t]+)?SECTION[ \t]+([A-Z]\d{2})\b[ \t]*[-–—:]?[ \t]*([A-Za-z][A-Za-z0-9 ,.&/()'\-]{2,60})?[ \t]*$", re.M)),
    ("verbatim", re.compile(r"^[ \t]*([A-Z]\d{2})[ \t]+[-–—:]?[ \t]*([A-Z][A-Z0-9 ,.&/()'\-]{2,60})[ \t]*$", re.M)),
    ("verbatim", re.compile(r"^[ \t]*(\d{1,2}\.)[ \t]+([A-Z][A-Z &/,'’\-]{3,60}?)[ \t]*$", re.M)),   # "1. SCOPE OF WORKS"
    ("verbatim", re.compile(r"^[ \t]*(Section[ \t]+\d{1,2})[ \t]*[-–—:][ \t]*([A-Z][A-Za-z0-9 ,.&/()'’\-]{2,60}?)[ \t]*$", re.M)),
    ("verbatim", re.compile(r"^[ \t]*((?:Ss|Pr|EF|PM|Ac)_\d{2}(?:_\d{2}){1,4})\b[ \t]*[-–—:]?[ \t]*([A-Za-z][A-Za-z0-9 ,.&/()'\-]{2,60})?[ \t]*$", re.M)),
]
ARTICLE = re.compile(r"^[ \t]*(\d\.\d{1,2})[ \t]+([A-Z][A-Z0-9 ,&/()'\-]{2,50})[ \t]*$", re.M)
ARTICLES_US = [ARTICLE]
# NBS clause numbers: "110 SCOPE OF WORK", "310 CABLE LADDER"
ARTICLES_UK = [re.compile(r"^[ \t]*(\d{1,2}\.\d{1,2})[ \t]+([A-Z][A-Za-z0-9 ,:&/()'’\-“”\"]{2,60}?)[ \t]*$", re.M),   # 1.5 Testing and Commissioning
               re.compile(r"^[ \t]*(\d{3})[ \t]+([A-Z][A-Z0-9 ,&/()'\-]{2,50})[ \t]*$", re.M)]
ARTICLES = ARTICLES_US
INDEX_EXTRA = None   # UK: NBS contents pages list many "Y61 ..." lines
STRICT = False       # UK: fixed sentence splitter + values only from the matched sentence (see configure())

def norm_line(line):
    return re.sub(r"\d+", "#", re.sub(r"\s+", " ", line.strip().lower()))

def blank_repeated_lines(pages):
    """Running headers/footers repeat on many pages; blank them (same length keeps offsets)."""
    if len(pages) < 6:
        return pages
    cnt = Counter()
    for t in pages:
        for ln in set(norm_line(l) for l in t.splitlines()):
            if len(ln) > 3:
                cnt[ln] += 1
    thresh = max(4, int(len(pages) * 0.25))
    rep = {ln for ln, c in cnt.items() if c >= thresh}
    out = []
    for t in pages:
        lines = t.split("\n")
        out.append("\n".join(" " * len(l) if norm_line(l) in rep else l for l in lines))
    return out

class Doc:
    def __init__(self, name, text):
        self.name = name
        raw_pages = normalise(text).split("\f")
        if raw_pages and not raw_pages[-1].strip():
            raw_pages = raw_pages[:-1]
        # footer section numbers must be read before blanking repeated lines
        self.footer_sec = []
        for t in raw_pages:
            m = FOOTER_SEC.search(t)
            self.footer_sec.append((m.group(1) or m.group(2)) if m else "")
        self.pages = blank_repeated_lines(raw_pages)
        self.index_pages = {i for i, t in enumerate(raw_pages)
                            if len(re.findall(r"\b\d{2} \d{2} \d{2}\b|\b0\d{4}\b", t)) >= 10
                            or re.search(r"TABLE\s+OF\s+CONTENTS|(?:INDEX|LIST)\s+OF\s+DRAWINGS", t)
                            or (INDEX_EXTRA is not None and INDEX_EXTRA(t))}
        self.empty_pages = sum(1 for t in raw_pages if len(t.strip()) < 40)
        self._index_sections()
        self._sent_cache = {}
        self._start_cache = {}

    def _index_sections(self):
        """Per page: sorted list of (offset, label). Carry the last label across pages."""
        self.headers = []
        self.carry_in = []
        carry = ""
        for pi_, t in enumerate(self.pages):
            self.carry_in.append(carry)
            hs = []
            if STRICT and pi_ in self.index_pages:   # a contents page lists every heading - don't carry them
                self.headers.append(hs)
                continue
            for kind, rx in HDR_PATTERNS:
                for m in rx.finditer(t):
                    num = m.group(1)
                    title = (m.group(2) or "").strip() if (m.lastindex or 0) >= 2 else ""
                    if kind == "exhibit":
                        label = f"Exhibit {num}"
                    elif kind == "division":
                        label = f"Division {num}" + (f" {title.title()}" if title else "")
                    elif kind == "roman":
                        label = f"Section {num}. {title.title()}"
                    elif kind == "verbatim":
                        label = num + (f" {title}" if title else "")
                    else:
                        label = num + (f" {title.title()}" if title else "")
                    hs.append((m.start(), "sec", label[:60]))
            for rx in ARTICLES:
                for m in rx.finditer(t):
                    hs.append((m.start(), "art", f"§{m.group(1)} {m.group(2).strip().title()}"[:40]))
            hs.sort()
            self.headers.append(hs)
            for off, kind, label in hs:
                if kind == "sec":
                    carry = label
        # article carry (reset when a new section starts)
        self.art_carry_in = []
        art = ""
        for hs in self.headers:
            self.art_carry_in.append(art)
            for off, kind, label in hs:
                art = label if kind == "art" else ("" if kind == "sec" else art)

    def section_at(self, pi, off):
        sec, art = "", ""
        for o, kind, label in self.headers[pi]:
            if o > off:
                break
            if kind == "sec":
                sec, art = label, ""
            else:
                art = label
        if not sec:
            sec = self.footer_sec[pi] or self.carry_in[pi]
            if not art:
                art = self.art_carry_in[pi]
        return (sec + (" " + art if art else "")).strip()

    def sentences(self, pi):
        """[(start, end, clean_text)] for page pi. Split on sentence ends, list markers, blank lines."""
        if pi in self._sent_cache:
            return self._sent_cache[pi]
        t = self.pages[pi]
        cuts = {0, len(t)}
        if STRICT:
            # a full stop at the end of a line followed by an un-indented line is still a sentence end
            # (the US splitter required indentation, so left-aligned paragraphs ran together and a
            # value from one clause could be reported against another)
            # (Word exports bullets as private-use Symbol-font characters such as U+F0B7)
            for m in re.finditer(r"(?<=[a-z0-9)\]\"'%])\.(?:[ \t]+|[ \t]*\n[ \t]*)(?=(?:[•\uf0a7\uf0b7\uf0d8\uf076\uf0fc][ \t]+)?[A-Z(])", t):
                cuts.add(m.end())
            for m in re.finditer(r"\n[ \t]*(?=\d{3}[ \t]+[A-Z]{2}|[-•–\uf0a7\uf0b7\uf0d8\uf076\uf0fc][ \t]+[A-Z])", t):   # NBS clauses, bullets
                cuts.add(m.end())
            for off, kind, label in self.headers[pi]:                          # headings stand alone
                cuts.add(off)
                eol = t.find("\n", off)
                cuts.add(eol if eol != -1 else len(t))
        else:
            for m in re.finditer(r"(?<=[a-z0-9)\]\"'%])\.[ \t]*\n?[ \t]+(?=[A-Z(])", t):
                cuts.add(m.end())
        for m in re.finditer(r"\n[ \t]*(?=(?:[A-Z]|\d{1,2}|[a-z])\.[ \t]+[A-Z\"(])", t):
            cuts.add(m.end())
        for m in re.finditer(r"\n[ \t]*(?=\d\.\d{1,2}[ \t]+[A-Z])", t):
            cuts.add(m.end())
        for m in re.finditer(r"\n[ \t]*\n", t):
            cuts.add(m.end())
        cuts = sorted(cuts)
        out = []
        for a, b in zip(cuts, cuts[1:]):
            s = clean(t[a:b])
            if len(s) >= 12:
                out.append((a, b, s))
        self._sent_cache[pi] = out
        self._start_cache[pi] = [x[0] for x in out]
        return out

    def _index(self, pi, off):
        sents = self.sentences(pi)
        return sents, bisect.bisect_right(self._start_cache[pi], off) - 1

    def context_at(self, pi, off, after=1):
        sents, i = self._index(pi, off)
        if i < 0:
            return ""
        return " ".join(x[2] for x in sents[i:i + 1 + after])

    def sentence_at(self, pi, off):
        sents, i = self._index(pi, off)
        if 0 <= i < len(sents):
            return sents[i]
        return (off, off, "")

def clean(s):
    return re.sub(r"\s+", " ", s.replace("\ufffd", "•")).strip()

def shorten(s, focus=None, limit=380):
    if len(s) <= limit:
        return s
    if focus:
        i = s.lower().find(focus.lower())
        if i > limit // 2:
            a = max(0, i - limit // 2)
            a = s.rfind(" ", 0, a) + 1
            frag = s[a:a + limit]
            return "…" + frag[:frag.rfind(" ")] + "…"
    return s[:s.rfind(" ", 0, limit)] + "…"

REF_ARTICLE = re.compile(r"Applicable\s+Publications|\bReferences?\b|Reference\s+Standards|Definitions|Related\s+(?:Requirements|Sections|Work|Documents)|Section\s+Includes|\bSummary\b", re.I)
STD_LINE = re.compile(r"\b[A-Z]{0,2}\d{1,5}[A-Z]?(?:\.\d+)?(?:/[A-Z]{0,2}\d+M?)?-\d{2}(?:\([A-Z]\d{4}\))?\s*[-–—]\s+[A-Z]|https?://|www\.")
REQ_VERB = re.compile(r"\b(?:shall|must|provide|use|apply|install|include|required|submit|furnish|perform)\b", re.I)

def not_a_requirement(d, pi, off, s):
    """TOC/drawing-index pages, reference/definition articles and standards lists are not requirements."""
    return pi in d.index_pages or bool(REF_ARTICLE.search(d.section_at(pi, off))) or bool(STD_LINE.search(s))

def iter_matches(docs, rx):
    """Yield (doc, page_index, match) for every match of rx on every page."""
    for d in docs:
        for pi, t in enumerate(d.pages):
            for m in rx.finditer(t):
                yield d, pi, m

def cite(d, pi, off):
    return {"Document": d.name, "Section": d.section_at(pi, off), "Page": pi + 1}

# =====================================================================================
# 2. Key-value extraction (dates, money, durations, percents)
# =====================================================================================
MONTHS = r"(?:January|February|March|April|May|June|July|August|September|October|November|December|Jan\.?|Feb\.?|Mar\.?|Apr\.?|Jun\.?|Jul\.?|Aug\.?|Sept?\.?|Oct\.?|Nov\.?|Dec\.?)"
KV = [
    re.compile(r"(?:(?:Mon|Tues|Wednes|Thurs|Fri|Satur|Sun)day,?\s+)?" + MONTHS + r"\s+\d{1,2},?\s+\d{4}(?:,?\s*(?:at\s+)?\d{1,2}(?::\d{2})?\s*[ap]\.?\s?m\.?)?", re.I),
    re.compile(r"\d{1,2}(?::\d{2})?\s*[ap]\.m\.", re.I),
    re.compile(r"\$\s?[\d,]+(?:\.\d{2})?\)?(?:\s*(?:per|for each|each)\s+(?:calendar\s+|working\s+)?(?:day|occurrence|week)(?:\s+per\s+day)?)?", re.I),
    re.compile(r"\d{1,3}(?:\.\d+)?\s*(?:%|percent)", re.I),
    re.compile(r"(?:\b[a-z\- ]{3,30}\s)?\(\d{1,4}\)\s+(?:consecutive\s+)?(?:calendar|working|business)\s+days|\b\d{1,4}\s+(?:consecutive\s+)?(?:calendar|working|business)\s+days", re.I),
    re.compile(r"\b\d{1,3}\s+days\b", re.I),
    re.compile(r"\b(?:Class|Classification)\s+[A-Z](?:-\d{1,2})?\b(?:\s+(?:or|and)\s+[A-Z](?:-\d{1,2})?)?"),
    re.compile(r"\"[A-Z](?:-\d{1,2})?\"\s+licen[cs]e(?:\s+or\s+an?\s+\"[A-Z](?:-\d{1,2})?\"\s+licen[cs]e)?", re.I),
]
KV_US = KV
KV_UK = [
    # 14 November 2026 / Friday 14th November 2026 (day-month order), then US order, then 14/11/2026
    re.compile(r"(?:(?:Mon|Tues|Wednes|Thurs|Fri|Satur|Sun)day,?\s+)?\d{1,2}(?:st|nd|rd|th)?\s+(?:of\s+)?" + MONTHS + r",?\s+\d{4}", re.I),
    KV_US[0],
    re.compile(r"\b\d{1,2}/\d{1,2}/(?:\d{4}|\d{2})\b"),
    re.compile(r"\b\d{1,2}:\d{2}(?:\s*(?:hrs|hours|noon|midday|[ap]\.?m\.?))?|\b\d{1,2}\.\d{2}\s*(?:hrs|hours|noon|midday|[ap]\.?m\.?)|\b12\s*noon\b|\bmidday\b|\b\d{1,2}\s*[ap]\.m\.", re.I),
    re.compile(r"[£$€]\s?[\d,]+(?:\.\d{2})?(?:\s*(?:million|m|k)\b)?(?:\s*(?:per|for\s+each|each|/)\s*(?:calendar\s+|working\s+)?(?:day|week|occurrence|month|claim|event)(?:\s+or\s+part\s+(?:thereof|of\s+a\s+week))?)?", re.I),
    re.compile(r"\d{1,3}(?:\.\d+)?\s*(?:%|percent|per\s+cent)", re.I),
    re.compile(r"\b(?:\d{1,2}|six|twelve|eighteen|twenty-four)[ \-]months?\b", re.I),
    re.compile(r"(?:\b[a-z\-]{3,20}\s)?\(\d{1,4}\)\s+(?:consecutive\s+)?(?:calendar\s+|working\s+|business\s+)?(?:days|weeks|months)|\b\d{1,4}\s+(?:consecutive\s+)?(?:calendar\s+|working\s+|business\s+)?(?:days|weeks|months)\b", re.I),
    # form of contract - kept case-sensitive so "option b" in prose is not read as NEC Option B
    re.compile(r"\bJCT\b[^.;\n]{0,60}?(?:Sub-?[Cc]ontract|Contract)(?:\s+20\d{2}|\s+\(\d{4}\))?|\bNEC\s?[34]?\s+(?:Engineering\s+(?:and|&)\s+Construction|Professional\s+Services?|Term\s+Service|Supply)\s+(?:Sub)?[Cc]ontract|\bNEC\s?[34]\b|\bFIDIC\b|\bOption\s+[A-F]\b|\bOption\s+X\d{1,2}\b"),
    re.compile(r"\b(?:CHAS|SSIP|Constructionline(?:\s+(?:Gold|Silver|Platinum))?|SafeContractor|NICEIC|NAPIT|ECA|CSCS|ECS|JIB|IPAF|PASMA|CPCS|SMSTS|SSSTS|ISO\s?(?:9001|14001|45001|27001))\b"),
]
KV_LIMIT = 3

def key_values(s, limit=3):
    found = []
    for rx in KV:
        for m in rx.finditer(s):
            v = clean(m.group(0)).strip(" ,.")
            v = re.sub(r"^(\$\s?[\d,]+(?:\.\d{2})?)\)", lambda mm: mm.group(1), v)
            v = re.sub(r"^(?:(?:shall|will|must|have|has|be|the|of|within|for|a|an|contractor|is|are)\s+)+", "", v, flags=re.I)
            if v and not any(v in f or f in v for f in found):
                found.append(v)
    return "; ".join(found[:max(limit, KV_LIMIT)])

# =====================================================================================
# 3. Bid-killers
# =====================================================================================
BID_KILLERS = OrderedDict([
    ("Bid due date / opening", [r"bids?\s+(?:will\s+be\s+|shall\s+be\s+)?(?:received|opened|due|submitted)[\s\S]{0,160}?(?:\d{1,2}:\d{2}|[ap]\.m\.|\d{4})", r"(?:proposal|offer)s?\s+(?:due|must\s+be\s+received)"]),
    ("Mandatory pre-bid / job-walk", [r"mandatory\s+pre-?(?:bid|proposal)", r"pre-?(?:bid|proposal)\s+(?:meeting|conference|site\s+visit)", r"job\s?walk", r"site\s+(?:visit|inspection)\s+(?:is\s+|will\s+be\s+)?(?:mandatory|required|scheduled)"]),
    ("RFI / questions deadline", [r"\b(?:questions?|inquiries|requests?\s+for\s+information|rfis?|rfi's)\b[\s\S]{0,100}?(?:deadline|no\s+later\s+than|received\s+by|due\s+(?:by|on)|within\s+\(?\w+\)?\s*(?:\(\d+\)\s*)?(?:business|calendar|working)?\s*days|prior\s+to\s+(?:the\s+)?bid)", r"deadline\s+for\s+(?:questions|inquiries|rfis?)"]),
    ("Substitution request deadline", [r"(?:requests?\s+for\s+substitution|substitution\s+requests?)[\s\S]{0,140}?(?:days|prior\s+to|before|no\s+later\s+than)"]),
    ("Bid bond / bid security", [r"security\s+shall\s+be\s+in\s+an\s+amount", r"bid\s+shall\s+not\s+be\s+considered\s+unless[\s\S]{0,80}security", r"bid\s+bond", r"bid\s+(?:guarantee|security|guaranty)", r"bidder'?s\s+(?:bond|guarant\w*|security)", r"proposal\s+(?:bond|guaranty)"]),
    ("Performance bond", [r"performance\s+bonds?"]),
    ("Payment bond", [r"payment\s+bonds?", r"labor\s+and\s+material\s+(?:payment\s+)?bonds?"]),
    ("Post-award bonds/insurance deadline", [r"forfeit(?:ure|ed)?[\s\S]{0,120}(?:guarant|bond|security)", r"within\s+\(?\w+\)?\s*(?:\(\d+\)\s*)?(?:calendar\s+|working\s+|business\s+)?days\s+(?:after|from|of)[\s\S]{0,60}award[\s\S]{0,120}(?:bond|insurance|contract)"]),
    ("Bid validity period", [r"hold\s+the\s+bidder\s+to\s+its\s+bid[\s\S]{0,60}days", r"bids?\s+(?:shall\s+)?(?:remain|be)\s+(?:valid|open|in\s+effect|firm)[\s\S]{0,60}days", r"(?:may\s+not|shall\s+not|no\s+bidder\s+(?:may|shall))\s+withdraw[\s\S]{0,80}days", r"acceptance\s+period"]),
    ("Liquidated damages", [r"liquidated\s+damages"]),
    ("Per-day penalties / fines", [r"\$\s?[\d,]+(?:\.\d\d)?\)?\s*per\s+(?:occurrence|calendar\s+day|working\s+day|day)"]),
    ("Contract time / completion", [r"\(?\d{1,4}\)?\s+(?:consecutive\s+)?(?:working|calendar)\s+days[\s\S]{0,160}?(?:to\s+complete|complet(?:e|ion)\s+(?:all|the|of\s+(?:all|the))\s+(?:work|project|contract)|shall\s+be\s+(?:substantially\s+)?completed)", r"(?:work|project)\s+shall\s+be\s+(?:substantially\s+)?complete(?:d)?\s+within\s+\(?\w+\)?", r"(?:complete|completion\s+of)\s+(?:all\s+)?(?:the\s+)?work[\s\S]{0,60}within\s+\(?\w+\)?\s*(?:\(\d+\)\s*)?(?:calendar|working)\s+days", r"contract\s+(?:time|period|duration)\s+(?:of|is|shall\s+be)\s+(?:\(?\d|\w+\s+\(\d)", r"substantial(?:ly)?\s+complet\w*[\s\S]{0,100}?(?:\(?\d{1,4}\)?\s+(?:consecutive\s+)?(?:calendar|working)\s+days|" + MONTHS + r"\s+\d{1,2},?\s+\d{4})"]),
    ("Insurance requirement", [r"(?:general|commercial\s+general|professional|auto(?:mobile)?|umbrella|excess)\s+(?:general\s+)?liability\s+insurance", r"builder'?s\s+risk", r"certificates?\s+of\s+insurance", r"bonds?\s+and\s+insurance", r"proof\s+of\s+insurance", r"insurance\s+(?:certificates?|requirements?|coverage)", r"workers'?\s+compensation\s+insurance"]),
    ("Prevailing wage / Davis-Bacon", [r"prevailing\s+wages?", r"davis-?\s?bacon", r"wage\s+(?:determination|decision|rates)"]),
    ("Required contractor license", [r"(?:must|shall)\s+(?:have|hold|possess)\s+(?:an?\s+)?\"?[A-Z](?:-\d{1,2})?\"?\s+licen[cs]e", r"\"[A-Z](?:-\d{1,2})?\"\s+licen[cs]e", r"(?:shall|must)\s+be\s+licensed\s+in\s+accordance", r"(?:contractor'?s?\s+)?licen[cs]e[\s\S]{0,80}?(?:Class\s+[A-Z](?:-\d{1,2})?|classification)", r"(?:shall|must)\s+(?:possess|hold|have)\s+(?:a\s+)?(?:valid\s+)?(?:\w+\s+)?(?:contractor'?s?\s+)?licen[cs]e"]),
    ("Retainage", [r"retainage", r"retention\s+(?:of|shall|will)", r"(?:withhold|retain)\s+\d{1,2}\s*(?:%|percent)"]),
    ("Alternates", [r"\badd(?:itive)?\s+alternates?", r"\bdeduct(?:ive)?\s+alternates?", r"\balternates?\s+(?:no\.?|#|bid|\d)", r"\bbid\s+alternates?"]),
    ("Addenda acknowledgement", [r"acknowledge?(?:ment|d)?\s+(?:of\s+)?(?:the\s+)?(?:receipt\s+of\s+)?(?:all\s+)?addend", r"addend(?:um|a)[\s\S]{0,80}(?:acknowledg|bid\s+form|proposal\s+form)", r"(?:check|monitor)\s+[\s\S]{0,60}addend"]),
])
BID_KILLER_RX = {k: [re.compile(p, re.I) for p in v] for k, v in BID_KILLERS.items()}
EXHIBIT_LIST = re.compile(r"(Exhibit|Attachment)\s+\"?([A-Z])\"?[ \t]*[:\-–]?[ \t]+([A-Z][A-Za-z0-9'&/,.\- ]{2,100}?)(?=\s+(?:Exhibit|Attachment)\s+\"?[A-Z]\b|\s{2,}|\n|$)")
FORM_RE = re.compile(r"\b(bid\s+form|proposal\s+form|bid\s+schedule|bidder'?s\s+proposal|schedule\s+of\s+bid\s+items)\b", re.I)
TOC_LINE = re.compile(r"(?:\.{4,}|\s{3,}\d{1,3}\s*$)|(?:Section\s+[IVXL]+\.[^.]{0,60}){2,}|(?:Exhibit\s+\"?[A-Z]\"?\s+\S[^.]{0,60}?){3,}", re.M)

BID_KILLERS_US = BID_KILLERS
# pattern kind drives hit_score() - US labels map to the same kinds so its scoring is unchanged
BK_KIND = {"Bid due date / opening": "due", "RFI / questions deadline": "queries", "Bid bond / bid security": "bond",
           "Performance bond": "bond", "Payment bond": "bond", "Post-award bonds/insurance deadline": "postaward",
           "Liquidated damages": "ld", "Per-day penalties / fines": "fines", "Required contractor license": "license"}

# UK tender / subcontract bid-killers. Labels are the item to check, not a claim about any contract.
BID_KILLERS_UK = OrderedDict([
    ("Tender return date / time", [r"tenders?\s+(?:shall|must|are\s+to|should|will)\s+be\s+(?:returned|submitted|received|delivered|uploaded|lodged)[\s\S]{0,160}?(?:\d{1,2}[:.]\d{2}|noon|midday|[ap]\.m\.|\d{4})", r"(?:tender\s+)?return\s+(?:date|deadline)", r"(?:deadline|closing\s+date)\s+for\s+(?:the\s+)?(?:receipt|return|submission)\s+of\s+tenders?", r"tender\s+(?:submission\s+)?deadline"]),
    ("Mandatory site visit / briefing", [r"(?:mid|pre)-?\s?tender\s+(?:site\s+)?(?:visit|meeting|briefing|interview)", r"site\s+visits?\s+(?:is\s+|are\s+|will\s+be\s+)?(?:mandatory|compulsory|required|arranged|scheduled)", r"(?:mandatory|compulsory)\s+(?:site\s+)?(?:visit|briefing|meeting)"]),
    ("Tender queries / clarifications deadline", [r"\b(?:tender\s+)?(?:queries|questions|clarifications?|enquiries|inquiries)\b[\s\S]{0,100}?(?:deadline|no\s+later\s+than|received\s+by|within\s+\(?\w+\)?\s*(?:\(\d+\)\s*)?(?:working\s+|calendar\s+|business\s+)?days|prior\s+to\s+(?:the\s+)?tender)", r"deadline\s+for\s+(?:queries|questions|clarifications|enquiries)"]),
    ("Tender validity period", [r"tenders?\s+(?:shall|must|will|should)\s+(?:remain|be)\s+(?:valid|open)[\s\S]{0,80}?(?:days|weeks|months)", r"(?:valid|open)\s+for\s+acceptance[\s\S]{0,40}?(?:days|weeks|months)", r"tender\s+validity", r"acceptance\s+period"]),
    ("Form of contract / pricing basis", [r"lump\s+sum\s+fixed\s+price|fixed\s+price\s+lump\s+sum|re-?measur\w+\s+(?:basis|contract)|(?:let|tendered|priced)\s+on\s+a\s+[^.\n]{0,30}?basis", r"\bJCT\b[\s\S]{0,80}?(?:Contract|Sub-?contract|20\d{2})", r"\bNEC\s?[34]?\b[\s\S]{0,60}?(?:Contract|Subcontract|Option)", r"\bFIDIC\b", r"(?:form|conditions)\s+of\s+(?:contract|sub-?contract)\s+(?:shall\s+be|will\s+be|is)"]),
    ("Contract amendments", [r"schedule\s+of\s+amendments", r"amendments\s+to\s+the\s+(?:JCT|NEC|contract|sub-?contract|conditions|standard\s+form)", r"bespoke\s+amendments", r"\bZ\s?clauses?\b|Option\s+Z\b", r"additional\s+conditions\s+of\s+contract"]),
    ("LADs / delay damages", [r"liquidated\s+(?:and\s+ascertained\s+)?damages", r"\bLADs?\b", r"delay\s+damages"]),
    ("Performance bond", [r"performance\s+bonds?", r"(?:contract|construction)\s+bond\b"]),
    ("Parent company guarantee", [r"parent\s+company\s+guarantee", r"\bPCG\b", r"holding\s+company\s+guarantee"]),
    ("Collateral warranties / third party rights", [r"collateral\s+warrant(?:y|ies)", r"third\s+party\s+rights"]),
    ("Retention", [r"retention\s+(?:of|shall|will|percentage|money|monies|bond|is|at)", r"\d{1,2}(?:\.\d{1,2})?\s*%\s+retention", r"retention\s+(?:\w+\s+){0,3}\d{1,2}(?:\.\d{1,2})?\s*%"]),
    ("Payment terms", [r"payment\s+(?:terms|cycle)", r"final\s+date\s+for\s+payment", r"pay\s?less\s+notice", r"interim\s+(?:valuations?|payments?|applications?)[\s\S]{0,80}?(?:monthly|days)"]),
    # only sentences that SET the period - "3 months prior to Sectional Completion" is a deadline for something else
    ("Contract period / completion", [r"(?:contract|construction)\s+(?:period|duration)\s+(?:of|is|shall\s+be|will\s+be)", r"date\s+for\s+(?:practical|sectional)\s+completion\s+(?:is|shall\s+be|will\s+be)", r"(?:practical|sectional)\s+completion\s+(?:date\s+)?(?:is|shall\s+be|will\s+be)\s+(?!\w+\s+prior)[^.\n]{0,40}?(?:\d{1,4}\s+(?:weeks|months)|\d{1,2}(?:st|nd|rd|th)?\s+" + MONTHS + r")", r"completion\s+date\s+(?:is|shall\s+be|will\s+be|of)", r"(?:complete|completed)\s+(?:the\s+works\s+)?within\s+\(?\d{1,3}\)?\s+weeks"]),
    ("Defects period / warranty", [r"defects?\s+(?:liability|rectification|notification)\s+period", r"rectification\s+period", r"\d{1,2}\s+months?\s+warranty|warranty\s+period\s+(?:of|is|shall)"]),
    ("Insurance requirement", [r"(?:public|employer'?s|products)\s+liability", r"professional\s+indemnity", r"contractors?'?\s+all\s+risks?", r"joint\s+names", r"insurance\s+(?:cover|requirements?|levels?|certificates?)"]),
    ("Design responsibility (CDP)", [r"contractor'?s\s+designed\s+portion", r"\bCDP\b", r"(?:sub-?)?contractor\s+(?:shall\s+be|is)\s+responsible\s+for\s+(?:the\s+)?(?:detailed\s+)?design", r"performance\s+specification", r"design\s+(?:and|&)\s+build"]),
    ("CDM duties", [r"\bCDM\b", r"principal\s+contractor", r"construction\s+phase\s+plan", r"pre-?construction\s+information"]),
    ("Accreditations / competence cards", [r"\b(?:CHAS|SSIP|Constructionline|SafeContractor|NICEIC|NAPIT|CSCS|ECS|JIB)\b", r"ISO\s?(?:9001|14001|45001)"]),
    ("Variants / qualified tenders", [r"variant\s+(?:bids?|tenders?)", r"alternative\s+(?:tenders?|bids?|offers?)", r"(?:qualified|conditional)\s+tenders?", r"qualifications\s+(?:to|in|within)\s+(?:the|your|their)\s+tender"]),
    ("Tender addenda acknowledgement", [r"tender\s+(?:addend(?:um|a)|bulletins?|clarification\s+notices?)", r"addend(?:um|a)[\s\S]{0,80}(?:acknowledg|form\s+of\s+tender)", r"acknowledge?\w*\s+(?:receipt\s+of\s+)?(?:all\s+)?(?:tender\s+)?(?:addend|bulletin)"]),
])
BK_KIND_UK = {"Tender return date / time": "due", "Tender queries / clarifications deadline": "queries",
              "Tender validity period": "validity", "Form of contract / pricing basis": "form", "LADs / delay damages": "ld",
              "Performance bond": "bond", "Parent company guarantee": "bond", "Retention": "pct",
              "Insurance requirement": "money", "Accreditations / competence cards": "accred",
              "Contract period / completion": "time", "Design responsibility (CDP)": "cdp", "Defects period / warranty": "time"}
FORM_RE_US = FORM_RE
FORM_RE_UK = re.compile(r"\b(with\s+the\s+tender\s+return|at\s+(?:bid|tender)\s+stage|in\s+(?:his|their|its|the)\s+tender\b|form\s+of\s+tender|tender\s+(?:return|submission)\s+(?:documents?|schedule|checklist)|pricing\s+(?:schedule|document)|activity\s+schedule|contract\s+sum\s+analysis|bills?\s+of\s+quantities|bid\s+form|bid\s+schedule)\b", re.I)
FORMS_LABEL = "Required bid forms / exhibits"

def hit_score(sentence, label):
    kind = BK_KIND.get(label, "")
    s = 0
    if re.search(r"[$£€]\s?[\d,]+", sentence): s += 3
    if re.search(r"\(\d+\)|\b\d+\s+(?:calendar|working|business)\s+days", sentence, re.I): s += 2
    if re.search(MONTHS + r"\s+\d{1,2}", sentence): s += 2
    if re.search(r"\d{1,2}:\d{2}|[ap]\.m\.", sentence, re.I): s += 1
    if re.search(r"\bshall\b|\bmust\b|\brequired\b|\bwill\b", sentence, re.I): s += 1
    if TOC_LINE.search(sentence) or len(sentence) < 30: s -= 4
    if re.search(r"non-?responsive|(?:will|shall)\s+be\s+rejected|shall\s+not\s+be\s+considered|disqualif", sentence, re.I): s += 4
    elif STRICT and re.search(r"non-?compliant|may\s+be\s+rejected|(?:will|shall|may)\s+(?:be\s+)?(?:excluded|disregarded)", sentence, re.I): s += 4
    if kind == "ld" and re.search(r"per\s+(?:calendar\s+)?(?:day|week)|each\s+(?:day|week)", sentence, re.I): s += 3
    if kind in ("bond", "postaward"):
        if re.search(r"\breturn", sentence, re.I): s -= 4
        if re.search(r"\d{1,3}\s*(?:%|percent)", sentence, re.I): s += 3
        if re.search(r"annul|forfeit", sentence, re.I) and kind == "postaward": s += 3
    if kind == "license" and re.search(r"\"[A-Z](?:-\d{1,2})?\"\s+licen|Class\s+[A-Z]\b", sentence): s += 4
    if kind == "fines" and re.search(r"delay|liquidated|not\s+completed", sentence, re.I): s -= 6
    if kind == "queries" and re.search(r"prior\s+to\s+(?:the\s+)?bid|before\s+(?:the\s+)?bid|bid\s+opening|questions", sentence, re.I): s += 3
    elif STRICT and kind == "queries" and re.search(r"(?:prior\s+to|before)\s+(?:the\s+)?tender|tender\s+return|clarification|queries", sentence, re.I): s += 3
    if STRICT:
        # UK: day-month dates, weeks/months, and the item-specific value the row exists for
        if re.search(r"\b\d{1,2}(?:st|nd|rd|th)?\s+" + MONTHS, sentence): s += 2
        if re.search(r"\b\d+\s+(?:weeks|months)\b", sentence, re.I): s += 2
        if kind == "validity" and re.search(r"\d+\)?\s+(?:days|weeks|months)", sentence, re.I): s += 3
        if kind == "form" and re.search(r"\bJCT\b|\bNEC\s?[34]?\b|\bFIDIC\b|lump\s+sum|fixed\s+price", sentence): s += 3
        if kind == "pct" and re.search(r"\d\s*%|per\s*cent", sentence, re.I): s += 3
        if kind == "cdp" and re.search(r"designed\s+portion|\bCDP\b|responsible\s+for\s+(?:the\s+)?(?:detailed\s+)?design", sentence, re.I): s += 3
        if kind == "accred": s += 2 * len(re.findall(r"\b(?:CHAS|SSIP|Constructionline|SafeContractor|NICEIC|NAPIT|CSCS|ECS|JIB|ISO\s?\d{4,5})\b", sentence))
    return s

def ld_note(s):
    c = s.lower()
    if re.search(r"per\s+(?:calendar\s+|working\s+)?(?:day|week)|each\s+(?:calendar\s+)?(?:day|week)|for\s+each\s+(?:calendar\s+)?(?:day|week)", c):
        return "per-day / per-week DELAY damages - carry in programme risk"
    if "execute the contract" in c or "guarant" in c or "bidder" in c:
        return "this is the BID-GUARANTEE forfeit, not a per-day delay LD - no per-day LD detected in this clause"
    return "read the clause for amount and trigger"

def _is_heading(s):
    """A heading / lead-in line (no closing full stop, or ends with ':') - its value sits in the next sentence."""
    return s.rstrip().endswith(":") or not re.search(r"[.;)]\s*$", s)

def _overlapping(rx, t):
    """Every match, restarting one character after each match start - so a match that is rejected for
    running into the next sentence does not swallow a valid match that starts inside it."""
    pos = 0
    while True:
        m = rx.search(t, pos)
        if not m:
            return
        yield m
        pos = m.start() + 1

def pass_bid_killers(docs):
    rows = []
    for label, rxs in BID_KILLER_RX.items():
        hits = []  # (score, doc, pi, off, sentence, match text, value borrowed from next sentence)
        seen = set()
        for d in docs:
            for pi, t in enumerate(d.pages):
                if pi in d.index_pages:
                    continue
                for rx in rxs:
                    for m in (_overlapping(rx, t) if STRICT else rx.finditer(t)):
                        a, b, s = d.sentence_at(pi, m.start())
                        if re.match(r"\W*(?:for\s+example|e\.g\.|such\s+as)", s, re.I):
                            continue
                        borrowed = False
                        if STRICT:
                            # score and value come from the matched sentence(s) only; the next sentence
                            # is read only when this one is a heading / lead-in, and the row says so
                            # a match that runs on into the next sentence pairs two unrelated clauses
                            # (e.g. "...at Practical Completion. The date ... is 26 weeks") - skip it; the
                            # later sentence gets its own chance to match
                            if d.sentence_at(pi, max(m.start(), m.end() - 1))[0] != a and not _is_heading(s):
                                continue
                            s_full = s
                            score = hit_score(s_full, label)
                            if not key_values(s_full) and _is_heading(s_full):
                                nxt = d.context_at(pi, m.start(), after=1)
                                if key_values(nxt) != key_values(s_full):
                                    s_full, borrowed = nxt, True
                        else:
                            # values (amount, date, days) often sit in the following sentence
                            s_full = d.context_at(pi, m.start(), after=1)
                            if len(s) < len(clean(m.group(0))):
                                s_full = clean(t[a:m.end() + 80])
                            score = hit_score(s_full, label)
                        key = (d.name, pi, s_full[:60].lower())
                        if key in seen:
                            continue
                        seen.add(key)
                        hits.append((score, d, pi, m.start(), s_full, clean(m.group(0)), borrowed))
        if not hits:
            rows.append({"Bid-killer": label, "Key value": "", "Clause": "NOT FOUND", "Document": "",
                         "Section": "", "Page": "", "Also on pages": "", "Confidence": "-",
                         "Notes": "not detected - confirm it truly does not apply before relying on it"})
            continue
        hits.sort(key=lambda h: (-h[0], h[1].name, h[2]))
        best = hits[0]
        sc, d, pi, off, s, mtxt, borrowed = best
        others = sorted({(h[1].name, h[2] + 1) for h in hits[1:] if (h[1].name, h[2]) != (d.name, pi)})
        also = ", ".join((f"{n} p{p}" if len(docs) > 1 else f"p{p}") for n, p in others[:15])
        if len(others) > 15:
            also += f" (+{len(others) - 15} more)"
        kind = BK_KIND.get(label, "")
        note = ld_note(s) if kind == "ld" else "verify the value against the page"
        if kind == "queries" and not re.search(r"bid|proposal|offer|tender", s, re.I):
            note = ("this looks like the construction-phase RFI / TQ process, not a tender-queries cutoff - check the ITT for one"
                    if STRICT else
                    "this is the construction-phase RFI process, not a pre-bid questions cutoff - check the bid invitation for one")
        if borrowed:
            note = "value is from the sentence after the heading - check it belongs to this item. " + note
        rows.append({"Bid-killer": label, "Key value": key_values(s), "Clause": shorten(s, mtxt),
                     **cite(d, pi, off), "Also on pages": also, "Confidence": "HIGH" if not borrowed else "MEDIUM",
                     "Notes": note})

    # required forms / exhibits register (+ attachments such as geotech reports)
    exhibits, attachments, first = OrderedDict(), OrderedDict(), None
    for d, pi, m in iter_matches(docs, EXHIBIT_LIST):
        kind, letter, title = m.group(1).lower(), m.group(2), clean(m.group(3))
        bucket = exhibits if kind == "exhibit" else attachments
        if letter not in bucket:
            bucket[letter] = title
            if kind == "exhibit":
                first = first or (d, pi, m.start())
    if exhibits:
        d, pi, off = first
        ex = sorted(exhibits)
        key = f"{len(ex)} exhibits ({ex[0]}–{ex[-1]})" if len(ex) > 1 else "1 exhibit"
        if attachments:
            key += f", {len(attachments)} attachments"
        listing = "; ".join(f"Exhibit {k}: {v}" for k, v in sorted(exhibits.items()))
        if attachments:
            listing += " | Attachments: " + "; ".join(f"{k}: {v}" for k, v in sorted(attachments.items()))
        rows.append({"Bid-killer": FORMS_LABEL,
                     "Key value": key,
                     "Clause": listing,
                     **cite(d, pi, off), "Also on pages": "", "Confidence": "HIGH",
                     "Notes": "complete every form - an incomplete bid form can get the proposal rejected"})
    else:
        hit = None
        forms = r"(?:bid\s+form|proposal\s+form|bid\s+schedule|bidder'?s\s+proposal" + (
            r"|with\s+the\s+tender\s+return|at\s+(?:bid|tender)\s+stage|in\s+(?:his|their|its|the)\s+tender\b|form\s+of\s+tender|pricing\s+(?:schedule|document)|activity\s+schedule|contract\s+sum\s+analysis|bills?\s+of\s+quantities|tender\s+(?:return|submission)\s+(?:documents?|schedule|checklist)" if STRICT else "") + ")"
        verbs = r"(?:submit|complete|fill(?:ed)?\s+(?:in|out)|enclose|attach|sign|return|issue|provide|supply|nominate|include)" if STRICT else r"(?:submit|complete|fill(?:ed)?\s+(?:in|out)|enclose|attach|sign)"
        more = []
        for d_, pi_, m_ in iter_matches(docs, FORM_RE):
            sent_ = d_.sentence_at(pi_, m_.start())[2]
            if re.search(verbs + r"\w*[^.]{0,60}" + forms + "|" + forms + r"[^.]{0,60}" + verbs, sent_, re.I) and not re.search(r"payment|compensation", sent_, re.I):
                if hit is None:
                    hit = (d_, pi_, m_)
                    if not STRICT:
                        break
                elif (d_.name, pi_) != (hit[0].name, hit[1]):
                    lbl = f"{d_.name} p{pi_ + 1}" if len(docs) > 1 else f"p{pi_ + 1}"
                    if lbl not in more:
                        more.append(lbl)
        if hit:
            d, pi, m = hit
            rows.append({"Bid-killer": FORMS_LABEL, "Key value": "",
                         "Clause": shorten(d.sentence_at(pi, m.start())[2], m.group(0)),
                         **cite(d, pi, m.start()), "Also on pages": ", ".join(more), "Confidence": "HIGH",
                         "Notes": "list and complete every required form" + (" - tender-stage returns are also asked for on the other pages listed" if more else "")})
        else:
            rows.append({"Bid-killer": FORMS_LABEL, "Key value": "", "Clause": "NOT FOUND",
                         "Document": "", "Section": "", "Page": "", "Also on pages": "", "Confidence": "-",
                         "Notes": "not detected - confirm it truly does not apply before relying on it"})
    return rows

# =====================================================================================
# 4. Hidden costs (included scope)
# =====================================================================================
HIDDEN = [
    ("No separate payment", r"no\s+separate\s+(?:payment|measurement|compensation|pay\s+item)", "carry this cost inside another bid item"),
    ("Incidental work", r"(?:considered|deemed)\s+(?:as\s+)?incidental|incidental\s+to\s+(?:the\s+)?(?:work|contract|bid|item|other)", "unpaid work - carry it in the related item"),
    ("Included in unit / lump-sum price", r"(?:included|include|includes)\s+(?:payment\s+)?(?:for\s+[^.]{0,40}\s+)?in\s+(?:the\s+)?(?:contract\s+)?(?:unit\s+prices?|lump\s+sum|contract\s+price|price\s+bid|bid\s+price|prices?\s+(?:paid|bid))", "price it into the named unit/lump-sum item"),
    ("At Contractor's expense", r"at\s+(?:the\s+)?contractor'?s\s+(?:own\s+)?(?:sole\s+)?(?:cost|expense)", "cost falls on you - price it"),
    ("No cost to Owner", r"at\s+no\s+(?:additional\s+|extra\s+)?(?:cost|charge|expense)\s+to\s+the\s+(?:owner|city|county|government|agency|state|district|authority|department)", "cost falls on you - price it"),
    ("No additional compensation", r"without\s+(?:additional|extra)\s+(?:compensation|cost|charge|payment)|no\s+additional\s+(?:compensation|payment|cost)\s+(?:will|shall)\s+be", "cost falls on you - price it"),
    ("Contractor bears cost", r"(?:shall\s+)?bear\s+(?:the|all)\s+(?:cost|costs|expense)|(?:cost|costs|expense)s?\s+(?:of|for)\s+[^.]{0,60}?(?:shall\s+be\s+)?(?:borne|paid)\s+by\s+the\s+contractor", "cost falls on you - price it"),
    ("Retesting at Contractor's cost", r"(?:retest(?:ing|s)?|re-test(?:ing|s)?|additional\s+tests?(?:ing)?)[^.]{0,100}(?:contractor|expense|cost)", "carry an allowance for failed-test retesting"),
]
HIDDEN_RX = [(c, re.compile(p, re.I), w) for c, p, w in HIDDEN]
HIDDEN_RX_US = HIDDEN_RX
_UK_TC = r"(?:(?:trade|sub-?|works|package|specialist)\s*)?contractor"
_UK_PAYER = r"(?:employer|client|owner|contractor|main\s+contractor|construction\s+manager|purchaser|authority|council|trust|end\s+user|tenant|landlord)"
HIDDEN_UK = [
    ("No separate payment", r"no\s+separate\s+(?:payment|measurement|compensation|pay\s+item|item)", "carry this cost inside another item"),
    ("Incidental work", r"(?:considered|deemed)\s+(?:as\s+)?incidental|incidental\s+to\s+(?:the\s+)?(?:works?|contract|item|other)", "unpaid work - carry it in the related item"),
    ("Included in Contract Sum / rates", r"(?:included|deemed\s+(?:to\s+be\s+)?included|allowed\s+for)\s+(?:with)?in\s+(?:the\s+|your\s+|his\s+|their\s+)?(?:(?:(?:trade|sub-?)\s*)?contract\s+sum|tender\s+(?:sum|price|figure)?|bid|offer|contract\s+price|lump\s+sum|(?:unit\s+)?rates|prices|preliminaries)|(?:contract|tender)\s+sum\s+(?:will|shall)\s+be\s+deemed\s+to\s+include", "price it into your rates / prelims"),
    ("Deemed to have allowed", r"deemed\s+to\s+have\s+(?:allowed|included|priced|made\s+(?:due\s+)?allowance|visited|inspected|satisfied)|(?:will|shall)\s+be\s+deemed\s+to\s+(?:be\s+)?included?\b", "you carry anything missed - price it or qualify it"),
    ("Unshown scope is yours", r"not\s+(?:currently\s+)?shown\s+(?:on|in)\s+(?:the\s+)?[^.]{0,60}?(?:drawings?|tender\s+information)|(?:whatever|irrespective\s+of)\s+(?:size|type)", "scope beyond the drawings - quantify it, or qualify the tender"),
    ("Variations / changes restricted", r"no\s+(?:consideration|reimbursement)\s+(?:will|shall)\s+be\s+made[^.]{0,60}variations?|(?:will\s+)?not\s+constitute\s+a\s+variation|no\s+adjustment\s+to\s+the\s+[^.]{0,30}?(?:contract\s+sum|price)|no\s+additional\s+(?:costs?|time|payment)\s+(?:or\s+(?:time|costs?)\s+)?(?:will|shall)\s+be\s+(?:considered|allowed|paid|granted)", "change risk sits with you - price it or qualify it"),
    ("Allow for", r"\b(?:sub-?)?contractors?\s+(?:shall|must|is\s+to|are\s+to|should)\s+(?:include\s+and\s+)?allow\s+for|\btenderers?\s+(?:shall|must|should|are\s+to)\s+allow\s+for|\ballow\s+for\s+all\b", "named cost with no pay item - price it"),
    ("At Contractor's expense", r"at\s+(?:the\s+)?" + _UK_TC + r"[’']?s?[’']?\s+(?:own\s+)?(?:sole\s+)?(?:cost|expense|risk\s+and\s+(?:cost|expense))|\b(?:solely|entirely)\s+at\s+(?:the\s+)?" + _UK_TC + r"[’']?s?[’']?\s+(?:own\s+)?(?:cost|expense|risk)|entirely\s+at\s+their\s+risk|(?:will|shall)\s+be\s+(?:solely|entirely)\s+provided\s+by\s+the\s+" + _UK_TC + r"|will\s+be\s+the\s+" + _UK_TC + r"[’']?s?[’']?\s+responsibility", "cost falls on you - price it"),
    ("No cost to Employer / Client", r"at\s+no\s+(?:additional\s+|extra\s+)?(?:cost|charge|expense)\s+to\s+the\s+" + _UK_PAYER, "cost falls on you - price it"),
    ("No additional payment", r"without\s+(?:additional|extra)\s+(?:compensation|cost|charge|payment)|no\s+(?:additional|extra)\s+(?:compensation|payment|cost|charge)\s+(?:will|shall)\s+be|\bno\s+extra\s+over\b|at\s+no\s+(?:extra|additional)\s+(?:cost|charge)\b|(?:repair|replac|rectif|remed|re-?test|attend|re-?offer)\w*[^.]{0,60}free\s+of\s+charge", "cost falls on you - price it"),
    ("Contractor bears cost", r"(?:shall\s+)?bear\s+(?:the|all)\s+(?:cost|costs|expense)|(?:cost|costs|expense)s?\s+(?:of|for)\s+[^.]{0,60}?(?:shall\s+be\s+)?(?:borne|paid)\s+by\s+the\s+" + _UK_TC + r"|" + _UK_TC + r"\s+(?:will|shall)\s+pay\s+(?:any|all)\s+costs|(?:shall|will)\s+be\s+chargeable\s+to\s+the\s+" + _UK_TC, "cost falls on you - price it"),
    ("Making good", r"\bmak(?:e|ing)\s+good\b", "making-good labour and materials - carry it"),
    ("Contra charges", r"contra[- ]?charge\w*|employ\s+others\s+to\s+do\s+so\s+and\s+charge|(?:costs?|expenses?)\s+(?:so\s+incurred\s+)?(?:will|shall|may)\s+be\s+(?:recovered|deducted|charged\s+back|set\s+off)", "deduction risk - price the obligation"),
    ("Retesting at Contractor's cost", r"(?:retest(?:ing|s)?|re-test(?:ing|s)?|additional\s+tests?(?:ing)?)[^.]{0,100}(?:contractor|expense|cost)", "carry an allowance for failed-test retesting"),
]
HIDDEN_RX_UK = [(c, re.compile(p, re.I), w) for c, p, w in HIDDEN_UK]

def pass_hidden(docs):
    rows, idx = [], {}
    for cat, rx, what in HIDDEN_RX:
        for d, pi, m in iter_matches(docs, rx):
            s = d.sentence_at(pi, m.start())[2]
            if not_a_requirement(d, pi, m.start(), s):
                continue
            key = re.sub(r"\W+", "", s.lower())[:120]
            if key in idx:
                r = idx[key]
                p = f"p{pi + 1}"
                if str(r["Page"]) != str(pi + 1) and p not in r["Also on pages"]:
                    r["Also on pages"] = (r["Also on pages"] + ", " + p).strip(", ")
                continue
            r = {"Hidden cost": cat, "Clause": shorten(s, m.group(0)), "What to carry": what,
                 **cite(d, pi, m.start()), "Also on pages": "", "Confidence": "HIGH"}
            idx[key] = r
            rows.append(r)
    rows.sort(key=lambda r: (r["Document"], r["Page"]))
    return rows or [{"Hidden cost": "NONE DETECTED", "Clause": "", "What to carry": "", "Document": "",
                     "Section": "", "Page": "", "Also on pages": "", "Confidence": "-"}]

# =====================================================================================
# 5. Sole-source / named products
# =====================================================================================
OR_EQUAL = re.compile(r"or\s+(?:an?\s+)?(?:approved\s+)?(?:equal|equivalent)|approved\s+equal", re.I)
NO_SUB = re.compile(r"no\s+substitut|substitutions?\s+(?:will|shall)\s+not\s+be\s+(?:permitted|accepted|allowed|considered)|(?:no|without)\s+(?:exception|equal)|sole\s+source|brand\s+name\s+(?:only|or\s+no)", re.I)
LIST_INTRO = re.compile(r"(?:acceptable|approved|available|qualified)\s+manufacturers?|provide\s+(?:products?\s+)?(?:by|from|of)\s+one\s+of\s+the\s+following|one\s+of\s+the\s+following\s+(?:manufacturers|products)|manufacturers?:\s*(?:subject\s+to|provide)|subject\s+to\s+compliance\s+with\s+requirements", re.I)
BOD = re.compile(r"basis[- ]of[- ]design", re.I)
GLOBAL_OR_EQUAL = re.compile(r"deemed\s+to\s+(?:read|be\s+followed\s+by)[\s\S]{0,80}?(?:or\s+equal|its\s+equal|their\s+equal)|or\s+its\s+equal\s+in\s+quality", re.I)
BY_MFR = re.compile(r"(?:([A-Z][\w®™\-]+(?:\s+[A-Z0-9][\w®™\-]*){0,3})\s*,?\s*(?:as\s+)?)?(?i:manufactured\s+by|made\s+by)\s+([A-Z][\w&.\-]*(?:\s+(?:&\s+)?[A-Z][\w&.\-]*){0,5})")
QUOTED = re.compile(r"\"([A-Z][^\"]{1,40})\"")
COMPANY = re.compile(r"^[ \t]*(?:\d{1,2}|[a-z])\.[ \t]+([A-Z][A-Za-z0-9&.,'\- ]{2,60}?(?:Inc\.?|Corp\.?|Co\.?|Company|LLC|Ltd\.?|Corporation|Industries|Manufacturing|Group|Products|Systems|International)?)[ \t]*[.;]?[ \t]*$", re.M)
OR_EQUAL_US, NO_SUB_US, GLOBAL_OR_EQUAL_US = OR_EQUAL, NO_SUB, GLOBAL_OR_EQUAL
OR_EQUAL_UK = re.compile(r"or\s+(?:an?\s+)?(?:approved\s+)?(?:equal|equivalent)|approved\s+equal|similar\s+approved|equivalent\s+approved|or\s+similar", re.I)
NO_SUB_UK = re.compile(NO_SUB_US.pattern + r"|substitutions?\s*:?\s+(?:is\s+|are\s+)?not\s+(?:permitted|accepted|allowed)|no\s+(?:alternatives?|equivalents?)\s+(?:will|shall)\s+be|(?:only|solely)\s+(?:the\s+)?(?:following|named)\s+(?:manufacturer|product)", re.I)
GLOBAL_OR_EQUAL_UK = re.compile(GLOBAL_OR_EQUAL_US.pattern + r"|(?:products?|items?|materials?)\s+(?:of\s+)?(?:an?\s+)?equivalent\s+(?:quality|standard|performance)[\s\S]{0,80}?(?:may|will)\s+be\s+(?:proposed|accepted|submitted|considered)|(?:where|if)\s+(?:products?|items?)\s+(?:are\s+)?(?:specified|named)\s+by\s+(?:proprietary|trade|manufacturer'?s?)\s+names?[\s\S]{0,120}?equivalent", re.I)
# UK: "Unistrut or approved equivalent" / "Legrand or similar approved"
NAMED_OR_EQUIV = re.compile(r"([A-Z][\w®™&.\-]*(?:[ \t]+[A-Z0-9][\w®™&.\-]*){0,3})[ \t]*(?:\((?!or\b)[^)\n]{0,30}\)[ \t]*)?,?[ \t]+\(?(?:or|/)[ \t]+(?:an?[ \t]+)?(?:approved[ \t]+)?(?:equivalent|equal|similar)\b")
# UK NBS product clauses: "Manufacturer: Unistrut Ltd" ... "Product reference: P1000"
NBS_MFR = re.compile(r"^[ \t]*[-•–]?[ \t]*(?:Recommended[ \t]+)?Manufacturer[ \t]*:[ \t]*(\S[^\n]{1,80})$", re.M | re.I)
NBS_MFR_OPEN = re.compile(r"^(?:Submit|Contractor'?s?\s+choice|To\s+be|Selected|Not\s+applicable|As\s|Any\b|TBC|TBA|N/?A\b|See\b|Refer)", re.I)
NOT_A_BRAND_UK = re.compile(r"^(?:BS|EN|ISO|IEC|BSI|CIBSE|BESA|IET|NHBC|NBS|Employer|Engineer|Contractor|Client|Architect|Supplier|Contract|Works?|Equipment|Cables?|Type|Class|Option|Grade|Standard|Specification|Drawing|Section|Clause|This|That|These|Other|Similar|Or)\b|^[\d\W]+$")
NOT_A_NAME = re.compile(r"^(?:Provide|Submit|Comply|Refer|Use|See|The|All|Each|Where|When|Product|Products|Material|Materials|Section|Part|Type|Class|Grade|Size|Color|Finish|Minimum|Maximum|Install|Manufacturer|Manufacturers|Basis|Subject|General|Description|Quality|Delivery|Storage|Warranty|Performance|Design|Source|Limitations)\b")

def pass_sole_source(docs):
    rows, seen = [], set()

    by_name = {}

    def add(cls, product, mfr, d, pi, off, s, conf):
        name = (product or mfr or "").lower()
        key = (cls.split(" ")[0], name)
        if key in by_name:
            r = by_name[key]
            pg = f"p{pi + 1}"
            if str(r["Page"]) != str(pi + 1) and pg not in r["Also on pages"]:
                r["Also on pages"] = (r["Also on pages"] + ", " + pg).strip(", ")
            if mfr and not r["Manufacturer(s)"]:
                r["Manufacturer(s)"] = mfr
            return
        r = {"Classification": cls, "Product": product or "", "Manufacturer(s)": mfr or "",
             "Clause": shorten(s, (mfr or product or "")[:20]), **cite(d, pi, off), "Also on pages": "", "Confidence": conf}
        by_name[key] = r
        rows.append(r)

    for d, pi, m in iter_matches(docs, GLOBAL_OR_EQUAL):
        s = d.sentence_at(pi, m.start())[2]
        add("GLOBAL OR-EQUAL CLAUSE - every trade name substitutable unless stated otherwise", "ALL NAMED PRODUCTS", "",
            d, pi, m.start(), s, "HIGH")
        break

    for d in docs:
        for pi, t in enumerate(d.pages):
            # (a) "X, manufactured by Y" / quoted trade names next to or-equal
            for m in BY_MFR.finditer(t):
                a, b, s = d.sentence_at(pi, m.start())
                prod, mfr = (m.group(1) or "").strip(), clean(m.group(2)).rstrip(".,;")
                if NOT_A_NAME.match(mfr) or len(mfr) < 3:
                    continue
                alt = re.match(r"\s*,?\s*or\s+(?!(?:an?\s+)?(?:approved\s+)?(?:equal|equivalent))([A-Z][\w&.\-]*(?:\s+[A-Z][\w&.\-]*){0,4})", t[m.end():m.end() + 80])
                if alt and not OR_EQUAL.search(s):
                    add("LIMITED LIST (listed manufacturers only)", prod, f"{mfr}; {clean(alt.group(1)).rstrip('.,;')}", d, pi, m.start(), s, "HIGH")
                    continue
                if OR_EQUAL.search(s):
                    add("OR-EQUAL (substitution allowed)", prod, mfr, d, pi, m.start(), s, "HIGH")
                elif NO_SUB.search(s):
                    add("SOLE SOURCE (no substitution)", prod, mfr, d, pi, m.start(), s, "HIGH")
                else:
                    add("NAMED - no substitution language (confirm)", prod, mfr, d, pi, m.start(), s, "MEDIUM")
            for m in QUOTED.finditer(t):
                a, b, s = d.sentence_at(pi, m.start())
                brand_before = re.search(r"[A-Z][\w&.\-]+\s*$", t[max(0, m.start() - 30):m.start()])
                if not (brand_before or re.search(r"manufactur|brand|model|trade\s+name|®|™", s, re.I)):
                    continue
                if OR_EQUAL.search(s) and not re.search(r"\"or\s+equal\"|\"approved\s+equal\"", m.group(0), re.I):
                    name = clean(m.group(1))
                    if not NOT_A_NAME.match(name) and not re.search(r"equal", name, re.I):
                        add("OR-EQUAL (substitution allowed)", name, "", d, pi, m.start(), s, "HIGH")
            # (b) manufacturer lists: "provide products by one of the following: 1. A 2. B"
            for m in LIST_INTRO.finditer(t):
                a, b, s = d.sentence_at(pi, m.start())
                block = t[m.end(): m.end() + 900]
                names = []
                for cm in COMPANY.finditer(block):
                    n = clean(cm.group(1)).rstrip(".,;")
                    if NOT_A_NAME.match(n) or len(n) < 3 or len(n.split()) > 7:
                        break
                    names.append(n)
                    if len(names) >= 8:
                        break
                if not names:
                    continue
                ctx = s + " " + clean(block[:400])
                if NO_SUB.search(ctx) or len(names) == 1 and not OR_EQUAL.search(ctx):
                    cls = "SOLE SOURCE (no substitution)" if NO_SUB.search(ctx) else "SINGLE NAMED MANUFACTURER (confirm)"
                elif OR_EQUAL.search(ctx):
                    cls = "OR-EQUAL (substitution allowed)"
                else:
                    cls = "LIMITED LIST (listed manufacturers only)"
                add(cls, "", "; ".join(names), d, pi, m.start(), s, "HIGH" if len(names) > 1 else "MEDIUM")
            # (c) basis of design
            for m in BOD.finditer(t):
                a, b, s = d.sentence_at(pi, m.start())
                nm = re.search(r"basis[- ]of[- ]design\s*(?:product)?\s*:?\s*([A-Z][\w&.\-]*(?:\s+[A-Z0-9][\w&.\-]*){0,5})", s, re.I)
                name = clean(nm.group(1)) if nm and not NOT_A_NAME.match(nm.group(1)) else ""
                add("BASIS OF DESIGN (equals via substitution procedure)", "", name, d, pi, m.start(), s, "MEDIUM" if not name else "HIGH")

            if not STRICT:
                continue
            # (d) UK: "X or approved equivalent" / "X or similar approved"
            for m in NAMED_OR_EQUIV.finditer(t):
                name = clean(m.group(1)).rstrip(".,;")
                if len(name) < 3 or NOT_A_NAME.match(name) or NOT_A_BRAND_UK.match(name):
                    continue
                a, b, s = d.sentence_at(pi, m.start())
                if not_a_requirement(d, pi, m.start(), s):
                    continue
                add("OR-EQUAL (substitution allowed)", name, "", d, pi, m.start(), s, "HIGH")
            # (e) UK: NBS "Manufacturer:" lines
            for m in NBS_MFR.finditer(t):
                val = clean(m.group(1)).rstrip(".,;")
                if NBS_MFR_OPEN.match(val) or len(val) < 2:
                    continue
                a, b, s = d.sentence_at(pi, m.start())
                # the product clause runs until a blank line, the next NBS clause or the next Manufacturer: line
                blk = t[m.start(): m.start() + 400]
                stop = re.search(r"\n[ \t]*\n|\n[ \t]*\d{3}[ \t]+[A-Z]|\n[ \t]*[-•–]?[ \t]*(?:Recommended[ \t]+)?Manufacturer[ \t]*:", blk[1:], re.I)
                ctx = clean(blk[: stop.start() + 1] if stop else blk)
                ref = re.search(r"Product\s+reference[ \t]*:[ \t]*([^\n]{1,60})", t[m.end(): m.end() + 300], re.I)
                prod = clean(ref.group(1)).rstrip(".,;") if ref and not NBS_MFR_OPEN.match(clean(ref.group(1))) else ""
                if NO_SUB.search(ctx):
                    cls, conf = "SOLE SOURCE (no substitution)", "HIGH"
                elif OR_EQUAL.search(ctx):
                    cls, conf = "OR-EQUAL (substitution allowed)", "HIGH"
                else:
                    cls, conf = "NAMED - no substitution language (confirm)", "MEDIUM"
                add(cls, prod, val, d, pi, m.start(), ctx, conf)

    order = ["GLOBAL", "SOLE", "SINGLE", "LIMITED", "NAMED", "BASIS", "OR-EQUAL"]
    rows.sort(key=lambda r: (next((i for i, k in enumerate(order) if r["Classification"].startswith(k)), 9), r["Document"], r["Page"]))
    return rows or [{"Classification": "NONE DETECTED", "Product": "", "Manufacturer(s)": "", "Clause": "",
                     "Document": "", "Section": "", "Page": "", "Also on pages": "", "Confidence": "-"}]

# =====================================================================================
# 6. Submittals
# =====================================================================================
SUB_TYPES = [
    ("Shop drawings", r"shop\s+drawings?|erection\s+drawings?|fabrication\s+drawings?|placing\s+drawings?|coordination\s+drawings?|layout\s+drawings?"),
    ("Product data", r"product\s+data|manufacturer'?s?\s+(?:literature|data|catalog\w*|cut\s*sheets?|technical\s+(?:data|information|literature)|product\s+information|specifications|installation\s+instructions|instructions|recommendations)|cut\s*sheets?|data\s+sheets?|catalog\s+data"),
    ("Samples", r"\bsamples?\b"),
    ("Mix design", r"mix\s+designs?|design\s+mix(?:es)?|mix\s+proportions"),
    ("Certificates", r"mill\s+(?:test\s+)?(?:certificates?|reports?|certs?)|certificat(?:es?|ions?)\s+(?:of\s+)?(?:compliance|conformance)|certif(?:ied|ication|icates?)"),
    ("Test reports", r"test\s+(?:reports?|results|data)|laboratory\s+reports?|inspection\s+reports?|testing\s+reports?"),
    ("Plans / procedures", r"(?:traffic\s+control|dewatering|shoring|safety|quality\s+control|hot\s+weather|cold\s+weather|erosion\s+control|stormwater|work|lift|lifting|demolition|disposal|staging|phasing|placement|pour|curing|protection|haul(?:ing)?|excavation|trenching|spill|emergency|site\s+logistics|waste\s+management)\s+plans?|\bSWPPP\b|method\s+statements?|\bprocedures?\b"),
    ("Schedule", r"(?:construction|progress|work|baseline|cpm|project)\s+schedules?|schedule\s+of\s+values"),
    ("Calculations / design", r"calculations?|design\s+data|delegated\s+design|(?:stamped|sealed)\s+by|engineered\s+(?:design|drawings?)"),
    ("Qualifications", r"qualifications?|resumes?|experience\s+(?:records?|lists?)|licen[cs]es?|references\s+of"),
    ("Warranty", r"warrant(?:y|ies)|guarantees?"),
    ("O&M / closeout", r"operation(?:s)?\s+and\s+maintenance|O\s?&\s?M|maintenance\s+(?:manuals?|data)|record\s+drawings?|as-?built|closeout|attic\s+stock|spare\s+parts"),
    ("Permits / insurance / payroll", r"permits?|insurance|certified\s+payroll|payrolls?"),
]
SUB_TYPE_RX = [(n, re.compile(p, re.I)) for n, p in SUB_TYPES]
SUB_TYPE_RX_US = SUB_TYPE_RX
SUB_TYPES_UK_EXTRA = {
    "Shop drawings": r"|installation\s+drawings?|builder'?s\s+work\s+drawings?|working\s+drawings?|technical\s+submissions?|containment\s+(?:layout\s+)?drawings?",
    "Plans / procedures": r"|\bRAMS\b|risk\s+assessments?|construction\s+phase\s+plan|commissioning\s+(?:plan|procedures?|method\s+statements?)",
    "Test reports": r"|electrical\s+installation\s+certificates?|\bEICs?\b|schedules?\s+of\s+test\s+results|commissioning\s+(?:records?|certificates?|results)",
    "O&M / closeout": r"|health\s+and\s+safety\s+file|handover\s+(?:documents?|information|pack)",
}
SUB_TYPE_RX_UK = [(n, re.compile(p + SUB_TYPES_UK_EXTRA.get(n, ""), re.I)) for n, p in SUB_TYPES]
SUBMIT_VERB_UK = re.compile(r"\bsubmi(?:t|ts|tted|tting|ssions?)\b|\b(?:furnish|issue[ds]?|provide[ds]?|produce[ds]?)\b[^.]{0,80}\b(?:for\s+(?:approval|review|comment|agreement)|to\s+the\s+(?:construction\s+manager|contract\s+administrator|project\s+manager|employer|client|engineer|architect|design\s+team))|\bsubmittals?\b\s*:", re.I)
SUBMIT_VERB = re.compile(r"\bsubmit(?:s|ted|ting)?\b|\bfurnish\b[^.]{0,60}\bfor\s+(?:approval|review)|\bsubmittals?\b\s*:", re.I)
SUB_EXCLUDE_UK_EXTRA = re.compile(r"\b(?:submitting\s+(?:the\s+|their\s+|its\s+)?tenders?|with\s+(?:the|their|its|your)\s+tender|as\s+part\s+of\s+(?:the|their|its|your)\s+tender|tender\s+(?:return|submission)|technical\s+queries|\bTQs?\b|compensation\s+events?|early\s+warnings?|payment\s+applications?|interim\s+applications?|applications?\s+for\s+payment|valuations?)\b", re.I)
SUB_EXCLUDE = re.compile(r"\b(?:submitting\s+bids?|bidder'?s?\s+(?:bond|guarant\w*)|in\s+lieu\s+of\s+depositing|claims?|arbitration|disputes?|grievances?|protests?|invoices?|applications?\s+for\s+payment|pay(?:ment)?\s+requests?|requests?\s+for\s+(?:additional|payment|extension|time)|written\s+notice|notice\s+(?:of|to)|as\s+part\s+of\s+the\s+bid|with\s+(?:the|their|its|his|her)\s+bid|bid\s+proceedings|subcontract(?:ing)?\s+(?:request|shall)|scheduling\s+request|rfi'?s?)\b", re.I)
SUB_NOISE = re.compile(r"\bthe\s+submitted\b|\bdiscuss\w*|^(?:[A-Z]\.\s+)?(?:submittal\s+procedures|submittals?\s+shall\s+(?:be\s+in\s+accordance\s+with|conform\s+to|comply\s+with))\s*:?\s*(?:the\s+(?:requirements|provisions)\s+of\s+)?section\s+[\d ]+[^.]{0,60}\.?$", re.I)
SAMPLING_NOT_SUBMITTAL = re.compile(r"take\s+samples|samples?\s+(?:will|shall)\s+be\s+(?:taken|secured|collected|obtained|tested)|sampling|stormwater\s+samples?|discharge\s+samples?|test\s+cylinders?|samples?\s+for\s+(?:slump|air|testing|tests)|obtaining\s+samples|collect\w*\s+samples", re.I)
TIMING = re.compile(r"within\s+\(?\w+\)?\s*(?:\(\d+\)\s*)?(?:calendar\s+|working\s+|business\s+)?days?[^.;,]{0,50}|(?:at\s+least|minimum\s+of|not\s+less\s+than|no\s+later\s+than)\s+\(?\w+\)?\s*(?:\(\d+\)\s*)?(?:calendar\s+|working\s+|business\s+)?(?:days?|weeks?)\s+(?:prior|before|in\s+advance|after)[^.;,]{0,40}|prior\s+to\s+(?:the\s+)?(?:start|beginning|commencement|fabrication|installation|placement|ordering|delivery|shipment|pouring|construction|work)[^.;,]{0,40}|before\s+(?:start|beginning|commencing|fabrication|installation|placement|ordering|delivery|shipment|pouring)[^.;,]{0,40}|\bweekly\b|\bmonthly\b|\bdaily\b|at\s+(?:the\s+)?pre-?construction\s+(?:meeting|conference)", re.I)
TIMING_UK = re.compile(r"\b(?:\d{1,3}|one|two|three|four|five|six|seven|eight|ten|twelve)\s+(?:working\s+|calendar\s+)?(?:days?|weeks?|months?)[’']?\s+(?:prior\s+to|before|after|in\s+advance\s+of|of\s+(?:the\s+)?(?:effective\s+start|appointment|commencement))[^.;,]{0,50}|within\s+(?:\d{1,3}|one|two|three|four|six|eight)\s+(?:working\s+)?(?:days?|weeks?|months?)\s+of\s+[^.;,]{0,40}|by\s+no\s+later\s+than\s+\w+day\s+each\s+week|\beach\s+(?:week|month)\b|\bweekly\b|\bmonthly\b", re.I)
SUB_LIST_HDR = re.compile(r"(?:list|schedule|summary)\s+of\s+(?:required\s+)?submittals", re.I)
SUB_LIST_ROW = re.compile(r"^[ \t]*(\d{1,2})\.[ \t]+(\S.{2,40}?)(?:[ \t]{2,}(\S.{2,50}?))?(?:[ \t]{2,}(\S.{2,50}?))?[ \t]*$", re.M)

def classify_submittal(s):
    types = [n for n, rx in SUB_TYPE_RX if rx.search(s)]
    return types

def pass_submittals(docs):
    rows, idx = [], {}
    # (a) the spec's own submittal schedule table, if it has one
    for d in docs:
        for pi, t in enumerate(d.pages):
            h = SUB_LIST_HDR.search(t)
            if not h:
                continue
            body = t[h.end():]
            for m in SUB_LIST_ROW.finditer(body):
                item = clean(m.group(2))
                ref = clean(m.group(3) or "")
                due = clean(m.group(4) or "")
                if not TIMING.search(due) and re.search(r"within|meeting|weekly|monthly|prior|days", ref, re.I):
                    ref, due = "", ref
                rows.append({"Type": "From the spec's submittal schedule", "Submittal": item,
                             "Timing": due, "Reference": ref, **cite(d, pi, h.start()),
                             "Also on pages": "", "Confidence": "HIGH"})
    listed = len(rows)
    table_pages = {(r["Document"], r["Page"]) for r in rows}
    # (b) every sentence that requires a submittal
    for d in docs:
        for pi in range(len(d.pages)):
            for a, b, s in d.sentences(pi):
                if (d.name, pi + 1) in table_pages or not_a_requirement(d, pi, a, s):
                    continue
                sec = d.section_at(pi, a)
                if re.search(r"bidder'?s\s+proposal|bid\s+form" + (r"|form\s+of\s+tender|instructions\s+to\s+tenderers" if STRICT else ""), sec, re.I):
                    continue
                in_sub_article = "submittal" in sec.lower()
                has_verb = bool((SUBMIT_VERB_UK if STRICT else SUBMIT_VERB).search(s))
                loose = False
                if STRICT and not has_verb:
                    # "will within 1 month of appointment provide a detailed programme" - a timed deliverable
                    loose = has_verb = bool(re.search(r"\b(?:shall|will|must|is\s+to|are\s+to)\s+(?:[\w\-]+\s+){0,6}?(?:provide|produce|issue|supply)\b", s, re.I) and TIMING_UK.search(s))
                if not (has_verb or in_sub_article):
                    continue
                types = classify_submittal(s)
                if not types and STRICT and has_verb and TIMING_UK.search(s):
                    types = ["Other (timed deliverable)"]
                if not types:
                    continue
                if STRICT and SUB_EXCLUDE_UK_EXTRA.search(s) and not re.search(r"shop\s+drawings?|product\s+data|samples?\s+for|installation\s+drawings?", s, re.I):
                    continue
                if SUB_EXCLUDE.search(s) and not re.search(r"shop\s+drawings?|product\s+data|samples?\s+for|mix\s+design|mill", s, re.I):
                    continue
                if types == ["Samples"] and SAMPLING_NOT_SUBMITTAL.search(s) and not re.search(r"submit[^.]{0,60}samples?|samples?\s+(?:for|of)\s+[^.]{0,30}(?:approval|review|selection)", s, re.I):
                    continue
                if TOC_LINE.search(s) or len(s) < 20 or SUB_NOISE.search(s):
                    continue
                if not has_verb and not re.match(r"^(?:[A-Z]|\d{1,2})\.\s", s) and len(s) < 40:
                    continue
                key = re.sub(r"\W+", "", s.lower())[:140]
                if key in idx:
                    r = idx[key]
                    p = f"p{pi + 1}"
                    if str(r["Page"]) != str(pi + 1) and p not in r["Also on pages"]:
                        r["Also on pages"] = (r["Also on pages"] + ", " + p).strip(", ")
                    continue
                tm = TIMING.search(s) or (TIMING_UK.search(s) if STRICT else None)
                r = {"Type": ", ".join(types[:2]), "Submittal": shorten(s), "Timing": clean(tm.group(0)) if tm else "",
                     "Reference": "", "Document": d.name, "Section": sec, "Page": pi + 1, "Also on pages": "",
                     "Confidence": "HIGH" if has_verb and not loose else "MEDIUM"}
                idx[key] = r
                rows.append(r)
    head, rest = rows[:listed], rows[listed:]
    rest.sort(key=lambda r: (r["Document"], r["Page"]))
    rows = head + rest
    return rows or [{"Type": "NONE DETECTED", "Submittal": "", "Timing": "", "Reference": "", "Document": "",
                     "Section": "", "Page": "", "Also on pages": "", "Confidence": "-"}]

# =====================================================================================
# 7. Costly items (trade watchlists)
# =====================================================================================
NEGATED = re.compile(r"\b(?:reimburs\w*|(?:paid|furnished|provided|borne)\s+by\s+the\s+(?:city|owner|government|county|agency)|prohibited|not\s+(?:be\s+)?(?:permitted|allowed|required|acceptable|used)|shall\s+not|do\s+not|will\s+not\s+be\s+(?:required|permitted)|is\s+not\s+required|are\s+not\s+required|no\s+\w+\s+(?:is|are)\s+(?:required|permitted))\b", re.I)
NEGATED_US = NEGATED
NEGATED_UK = re.compile(r"\b(?:reimburs\w*|(?:paid|furnished|provided|borne|supplied)\s+by\s+(?:the\s+)?(?:city|owner|government|county|agency|employer|client|others|main\s+contractor)|by\s+others|prohibited|not\s+(?:be\s+)?(?:permitted|allowed|required|acceptable|used)|shall\s+not|do\s+not|will\s+not\s+be\s+(?:required|permitted)|is\s+not\s+required|are\s+not\s+required|no\s+\w+\s+(?:is|are)\s+(?:required|permitted))\b", re.I)

def compile_term(pat):
    return re.compile(r"(?<![A-Za-z0-9])(?:" + pat.replace(" ", r"[\s\-]+") + r")(?![A-Za-z0-9])", re.I)

def pass_costly(docs, trades, watch):
    rows = []
    for tr in trades:
        for entry in watch.get(tr, []):
            rx = compile_term(entry["match"])
            pages, first, negated_only = [], None, True
            for d, pi, m in iter_matches(docs, rx):
                s = d.sentence_at(pi, m.start())[2]
                if TOC_LINE.search(s) or len(s) < 25 or not_a_requirement(d, pi, m.start(), s):
                    continue
                label = (f"{d.name} p{pi + 1}" if len(docs) > 1 else f"p{pi + 1}")
                if label not in pages:
                    pages.append(label)
                neg = bool(NEGATED.search(s))
                rank = (neg, not REQ_VERB.search(s))
                if first is None or rank < first[6]:
                    first = (d, pi, m.start(), s, neg, m.group(0), rank)
                if not neg:
                    negated_only = False
            if not first:
                continue
            d, pi, off, s, neg, mt, _ = first
            rows.append({"Trade": tr, "Costly item": entry["item"], "Why it costs": entry["why"],
                         "Evidence": shorten(s, mt, 300), **cite(d, pi, off),
                         "All pages": ", ".join(pages[:25]) + (f" (+{len(pages) - 25})" if len(pages) > 25 else ""),
                         "Mentions": len(pages), "Confidence": "MEDIUM",
                         "Notes": "stated as prohibited / not required / reimbursed - check before pricing" if negated_only else ""})
    return rows or [{"Trade": "", "Costly item": "NONE DETECTED", "Why it costs": "", "Evidence": "", "Document": "",
                     "Section": "", "Page": "", "All pages": "", "Mentions": 0, "Confidence": "-", "Notes": ""}]

# =====================================================================================
# 8. RFI list (price-relevant ambiguity only)
# =====================================================================================
AUTH_US = r"(?:engineer|architect|owner|city|county|director|inspector|contracting\s+officer|\bCOR\b|resident\s+engineer|government|agency|district|consultant)"
AUTH_UK = AUTH_US[:-1] + r"|employer|client|contract\s+administrator|\bCA\b|project\s+manager|supervisor|principal\s+designer|main\s+contractor|clerk\s+of\s+works|quantity\s+surveyor|\bQS\b|services\s+engineer|M&E\s+consultant|construction\s+manager|validation\s+(?:manager|engineer)|design\s+team|designer)"
def rfi_rules(AUTH, uk=False):
    rules = [
        ("Unforeseen condition - who pays?", 3, r"(?:rock|ledge|groundwater|ground\s+water|contaminat\w*|hazardous|unsuitable\s+(?:material|soil)s?|obstructions?|unknown\s+utilities|buried\s+\w+)[^.]{0,80}?(?:if|where|when|as)\s+encountered|(?:if|where|when)\s+encountered",
         "Is this paid as extra work or included in the bid? Please give the pay item / unit price basis."),
        ("Risk assumed by bidder", 3, r"(?:should|shall)\s+be\s+presumed|(?:contractor|bidders?)\s+(?:shall|should|must)\s+assume\s+(?:that|the\s+(?:existing|site|soil|ground|presence|condition))|it\s+shall\s+be\s+assumed|no\s+claims?\s+(?:will|shall)\s+be\s+(?:allowed|considered)",
         "Please confirm the assumption to price, or provide the data (e.g. geotech) so it can be priced rather than guessed."),
        ("Scope set later by the Engineer", 3, r"as\s+(?:may\s+be\s+)?directed(?:\s+by\s+(?:the\s+)?" + AUTH + r")?|(?:as|if\s+and\s+as|when|where)\s+(?:required|requested|deemed\s+necessary|determined|ordered)\s+by\s+(?:the\s+)?" + AUTH,
         "Please define the extent / quantity so it can be priced. As written, scope is decided later."),
        ("Subjective determination", 2, r"in\s+the\s+opinion\s+of\s+(?:the\s+)?" + AUTH + r"|" + AUTH + r"'?s?\s+(?:sole\s+)?discretion|(?:deemed|considered)\s+(?:necessary|unsuitable|unacceptable|defective)\s+by",
         "Please state the objective criteria that will be used for this determination."),
        ("Undefined acceptance criteria", 2, r"to\s+the\s+" + (r"(?:entire\s+|full\s+|reasonable\s+)?" if uk else "") + r"satisfaction\s+of|satisfactory\s+to\s+the\s+" + AUTH + r"|acceptable\s+to\s+the\s+" + AUTH,
         "Please give measurable acceptance criteria (tolerance, test, standard)."),
        ("Match existing", 2, r"match(?:ing)?\s+(?:the\s+)?existing|(?:same\s+as|to\s+match)\s+(?:the\s+)?(?:existing|adjacent)",
         "Please identify the existing product / provide a sample or spec to match."),
        ("Undefined standard", 1, r"(?:industry|trade)\s+standards?|good\s+(?:engineering|trade|construction)\s+practice|first[- ]class\s+(?:workmanship|manner|condition)|workmanlike\s+manner|best\s+(?:industry\s+)?practices?|highest\s+(?:quality|standard)",
         "Please cite the specific standard / edition that governs."),
    ]
    if uk:
        rules += [
            ("Provisional / PC sum - scope undefined", 2, r"provisional\s+sums?|prime\s+cost\s+sums?|\bPC\s+sums?\b",
             "Please confirm whether this is a defined or undefined provisional sum and exactly what it covers, so programme and preliminaries can be priced."),
            ("Approval-based acceptance", 2, r"(?:to|subject\s+to)\s+the\s+(?:prior\s+)?approval\s+of\s+(?:the\s+)?" + AUTH,
             "Please state the criteria the approval will be judged against, and the review period, so it can be programmed and priced."),
        ]
    return rules
RFI_RULES = rfi_rules(AUTH_US)
RFI_RX = [(c, w, re.compile(p, re.I), q) for c, w, p, q in RFI_RULES]
RFI_RX_US = RFI_RX
RFI_RX_UK = [(c, w, re.compile(p, re.I), q) for c, w, p, q in rfi_rules(AUTH_UK, uk=True)]
RFI_EXCLUDE = re.compile(r"\b(?:payments?|retainage|retention|indemnif\w*|surety|bonds?|bonding|insurance|insurer|rating|execute\s+the\s+contract|award\w*|counsel|suspension|suspend|terminat\w*|council|patent\w*|royalt\w*|progress\s+estimate|invoice\w*|records?|documents?|copies|personnel|instructors?|superintendent|staff|employees?)\b", re.I)
SCOPE_VERB = re.compile(r"\b(?:remov\w*|dispos\w*|furnish\w*|install\w*|provid\w*|replac\w*|repair\w*|excavat\w*|plac\w*|construct(?:ed)?\b|relocat\w*|protect\w*|clean\w*|haul\w*|backfill\w*|shor\w*|dewater\w*|test\w*|restor\w*|maintain\w*|erect\w*|patch\w*|resurfac\w*|stripe|striping|cur(?:e|ed|ing))\b", re.I)

SCOPE_VERB_UK = re.compile(SCOPE_VERB.pattern[:-3] + r"|mak(?:e|ing)\s+good|terminat(?:e|ed|ing|ion)s?\s+(?:of\s+)?(?:cables?|conductors?)|design\w*|demonstrat\w*|obtain\w*|commission\w*|waive\w*)\b", re.I)

def pass_rfi(docs):
    rows, idx = [], {}
    for cat, weight, rx, q in RFI_RX:
        for d, pi, m in iter_matches(docs, rx):
            s = d.sentence_at(pi, m.start())[2]
            if TOC_LINE.search(s) or len(s) < 25 or RFI_EXCLUDE.search(s) or not_a_requirement(d, pi, m.start(), s):
                continue
            if cat in ("Scope set later by the Engineer", "Undefined acceptance criteria", "Subjective determination", "Approval-based acceptance") and not (SCOPE_VERB_UK if STRICT else SCOPE_VERB).search(s):
                continue
            key = re.sub(r"\W+", "", s.lower())[:120]
            if key in idx:
                r = idx[key]
                p = f"p{pi + 1}"
                if str(r["Page"]) != str(pi + 1) and p not in r["Also on pages"]:
                    r["Also on pages"] = (r["Also on pages"] + ", " + p).strip(", ")
                continue
            c = cite(d, pi, m.start())
            quote = shorten(s, m.group(0), 220)
            ref = f"{c['Section']}, PDF p.{c['Page']}".strip(", ")
            r = {"Risk": cat, "Rank": weight, "Ambiguous clause": shorten(s, m.group(0)),
                 "Draft RFI": f"Ref. {ref}: \"{quote}\" - {q}", **c, "Also on pages": "", "Confidence": "MEDIUM"}
            idx[key] = r
            rows.append(r)
    rows.sort(key=lambda r: (-r["Rank"], r["Document"], r["Page"]))
    for r in rows:
        r["Rank"] = {3: "High", 2: "Medium", 1: "Low"}[r["Rank"]]
    return rows or [{"Risk": "NONE DETECTED", "Rank": "", "Ambiguous clause": "", "Draft RFI": "", "Document": "",
                     "Section": "", "Page": "", "Also on pages": "", "Confidence": "-"}]

# =====================================================================================
# 9. Excel
# =====================================================================================
WIDTH = {"Clause": 70, "Submittal": 70, "Ambiguous clause": 60, "Draft RFI": 70, "Evidence": 60,
         "Bid-killer": 30, "Hidden cost": 26, "Classification": 34, "Costly item": 28, "Why it costs": 34,
         "What to carry": 30, "Product": 22, "Manufacturer(s)": 34, "Type": 22, "Timing": 26, "Reference": 24,
         "Key value": 30, "Section": 28, "Document": 22, "Page": 7, "Mentions": 9, "Confidence": 11, "Rank": 8,
         "Risk": 26, "Trade": 10, "Also on pages": 22, "All pages": 26, "Notes": 34, "Item": 40, "Detail": 90}
# "Page" is the PDF page (1 = first sheet of the file), not the printed page number - say so on the sheet
HEADER_NAME = {"Page": "PDF page"}

def write_xlsx(path, sheets, multi_doc):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    hdr_fill = PatternFill("solid", fgColor="16263C")
    hdr_font = Font(bold=True, color="FFFFFF")
    fills = {"HIGH": "E0EEE2", "MEDIUM": "F6E7CF", "-": "EFEDE8", "High": "F8D7D3", "Medium": "F6E7CF", "Low": "E5EDF5"}
    for name, rows, kind in sheets:
        ws = wb.create_sheet(name[:31])
        if kind == "summary":
            ws.column_dimensions["A"].width = 44
            ws.column_dimensions["B"].width = 40
            ws.column_dimensions["C"].width = 18
            ws.column_dimensions["D"].width = 60
            for r in rows:
                ws.append(r["cells"])
                row = ws.max_row
                if r.get("style") == "title":
                    ws.cell(row, 1).font = Font(bold=True, size=14)
                elif r.get("style") == "h":
                    for c in range(1, 5):
                        ws.cell(row, c).fill = hdr_fill
                        ws.cell(row, c).font = hdr_font
                elif r.get("style") == "miss":
                    ws.cell(row, 1).font = Font(color="9A3B2E")
                for c in range(1, 5):
                    ws.cell(row, c).alignment = Alignment(wrap_text=True, vertical="top")
            continue
        cols = [c for c in rows[0].keys() if multi_doc or c != "Document"]
        ws.append([HEADER_NAME.get(c, c) for c in cols])
        for c in range(1, len(cols) + 1):
            ws.cell(1, c).fill = hdr_fill
            ws.cell(1, c).font = hdr_font
        for r in rows:
            ws.append([r.get(c, "") for c in cols])
            for key in ("Confidence", "Rank"):
                if key in cols:
                    v = str(r.get(key, ""))
                    if v in fills:
                        ws.cell(ws.max_row, cols.index(key) + 1).fill = PatternFill("solid", fgColor=fills[v])
            if "Key value" in cols and r.get("Key value"):
                ws.cell(ws.max_row, cols.index("Key value") + 1).font = Font(bold=True)
            if r.get("Clause") == "NOT FOUND":
                ws.cell(ws.max_row, cols.index("Clause") + 1).font = Font(color="9A3B2E", bold=True)
        for i, col in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = WIDTH.get(col, 16)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    try:
        wb.save(path)
    except PermissionError:
        sys.exit(f"Cannot write {path} - it is probably open in Excel. Close it (or pass a different --out) and run again.")

def build_summary(docs, bk, hid, ss, sub, cost, rfi, trades):
    R = []
    add = lambda *cells, style=None: R.append({"cells": list(cells), "style": style})
    multi = len(docs) > 1
    where = lambda r: (f"{r['Document']} p{r['Page']}" if multi else f"p{r['Page']}")
    add("SPEC BID-REVIEW - read this first", style="title")
    names = ", ".join(d.name for d in docs)
    pages = sum(len(d.pages) for d in docs)
    add(f"Documents: {names}", f"{pages} pages", "", f"Region: {REGION.upper()}. Trade watchlists: {', '.join(trades)}")
    add("Page numbers (p12) are PDF page numbers - 1 = first sheet of the PDF - not the printed page number.")
    found = [r for r in bk if r["Clause"] != "NOT FOUND"]
    div00 = len(found) >= 3
    if not div00:
        add(*TECH_ONLY_WARNING, style="miss")
    add("")
    add("BID-KILLERS", "Key value", "Where", "Clause", style="h")
    for r in found:
        add(r["Bid-killer"], r["Key value"], f"{r['Section'][:30]} {where(r)}".strip(), r["Clause"][:220])
    missing = [r["Bid-killer"] for r in bk if r["Clause"] == "NOT FOUND"]
    if missing:
        add("Not found in this document:", "", "", "; ".join(missing), style="miss")
    add("")
    real_hid = [r for r in hid if r["Hidden cost"] != "NONE DETECTED"]
    add(f"HIDDEN COSTS ({len(real_hid)}) - scope you pay for without a pay item", "What to carry", "Where", "Clause", style="h")
    for r in real_hid[:12]:
        add(r["Hidden cost"], r["What to carry"], where(r), r["Clause"][:220])
    if len(real_hid) > 12:
        add(f"... {len(real_hid) - 12} more on the Hidden Costs tab")
    add("")
    add("NAMED PRODUCTS", "Count", "", "", style="h")
    for cls, n in Counter(r["Classification"] for r in ss if r["Classification"] != "NONE DETECTED").most_common():
        add(cls, n)
    add("")
    real_sub = [r for r in sub if r["Type"] != "NONE DETECTED"]
    add(f"SUBMITTALS ({len(real_sub)})", "Count", "", "", style="h")
    type_counts = Counter()
    for r in real_sub:
        for t in r["Type"].split(", "):
            type_counts[t] += 1
    for t, n in type_counts.most_common():
        add(t, n)
    add("")
    real_cost = [r for r in cost if r["Costly item"] != "NONE DETECTED"]
    add(f"COSTLY ITEMS ({len(real_cost)})", "Why it costs", "Pages", "", style="h")
    for r in sorted(real_cost, key=lambda r: -r["Mentions"])[:15]:
        add(f"[{r['Trade']}] {r['Costly item']}", r["Why it costs"], r["All pages"][:40], r["Notes"])
    add("")
    real_rfi = [r for r in rfi if r["Risk"] != "NONE DETECTED"]
    add(f"TOP RFIs TO SEND ({len(real_rfi)} total)", "Risk", "Where", "Draft RFI", style="h")
    for r in real_rfi[:8]:
        add(r["Risk"], r["Rank"], where(r), r["Draft RFI"][:300])
    return R

# =====================================================================================
# 10. Region
# =====================================================================================
REGION = "us"
TECH_ONLY_WARNING = ("⚠ Little or no Division 00/01 content detected.", "", "",
    "This looks like technical specifications only. Bid-killers (bid bond, LDs, wage rates, bid due date) usually live in the bid invitation / solicitation (e.g. SF-1442, Instructions to Bidders). Run that document too: py scripts/bid_review.py spec.pdf solicitation.pdf")
TRADES = {"us": ["concrete", "mep", "finishes", "universal"],
          "uk": ["concrete", "mep", "electrical", "finishes", "universal_uk"]}
UNIVERSAL = {"us": "universal", "uk": "universal_uk"}

def configure(region):
    """Swap the pattern tables. 'us' is the original engine exactly; 'uk' adds UK wording on top."""
    global REGION, STRICT, HDR_PATTERNS, ARTICLES, INDEX_EXTRA, KV, KV_LIMIT, BID_KILLER_RX, BK_KIND, FORM_RE
    global FORMS_LABEL, NEGATED, HIDDEN_RX, OR_EQUAL, NO_SUB, GLOBAL_OR_EQUAL, SUB_TYPE_RX, RFI_RX, TECH_ONLY_WARNING
    REGION = region
    if region == "us":
        return
    STRICT = True
    HDR_PATTERNS, ARTICLES = HDR_PATTERNS_UK, ARTICLES_UK
    INDEX_EXTRA = lambda t: (len(re.findall(r"(?m)^[ \t]*[A-Z]\d{2}[ \t]+\S", t)) >= 10
                             or len(re.findall(r"(?m)^[ \t]*Section[ \t]+\d{1,2}[ \t]*[-–—]", t)) >= 6
                             or bool(re.search(r"(?m)^[ \t]*CONTENTS[ \t]*$", t)))
    KV, KV_LIMIT = KV_UK, 4
    BID_KILLER_RX = {k: [re.compile(p, re.I) for p in v] for k, v in BID_KILLERS_UK.items()}
    BK_KIND = BK_KIND_UK
    FORM_RE, FORMS_LABEL = FORM_RE_UK, "Required tender returns / forms"
    HIDDEN_RX = HIDDEN_RX_UK
    OR_EQUAL, NO_SUB, GLOBAL_OR_EQUAL = OR_EQUAL_UK, NO_SUB_UK, GLOBAL_OR_EQUAL_UK
    SUB_TYPE_RX = SUB_TYPE_RX_UK
    NEGATED = NEGATED_UK
    RFI_RX = RFI_RX_UK
    TECH_ONLY_WARNING = ("⚠ Little or no tender / contract content detected.", "", "",
        "This looks like a technical spec only (e.g. NBS work sections). Tender-killers - return date, LADs, bonds, "
        "retention, form of contract, amendments - usually live in the Invitation to Tender / Instructions to Tenderers / "
        "Contract Particulars / subcontract enquiry. Run those too: py scripts/bid_review.py spec.pdf itt.pdf")

# =====================================================================================
# 11. Main
# =====================================================================================
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("inputs", nargs="+", help="spec PDF(s) or pdftotext -layout .txt file(s)")
    ap.add_argument("--region", default="uk", choices=["uk", "us"], help="uk (default): UK tenders, NBS, JCT/NEC. us: CSI / US bid documents")
    ap.add_argument("--trade", default="all", help="all | electrical | mep | concrete | finishes (comma-separated ok)")
    ap.add_argument("--out", default=str(ROOT / "outputs" / "spec-review.xlsx"))
    args = ap.parse_args()

    # validate cheap things before the slow extraction
    configure(args.region)
    watch = json.loads(CONFIG.read_text(encoding="utf-8"))
    if args.trade == "all":
        trades = TRADES[args.region]
    else:
        trades = list(OrderedDict.fromkeys([t.strip() for t in args.trade.split(",") if t.strip()] + [UNIVERSAL[args.region]]))
        for t in trades:
            if t not in watch or t.startswith("_"):
                sys.exit(f"Unknown trade '{t}'. Choose from: all, {', '.join(k for k in watch if not k.startswith('_') and not k.startswith('universal'))}")
    for inp in args.inputs:
        if not Path(inp).exists():
            sys.exit(f"Input not found: {inp}")

    exe = find_pdftotext() if any(Path(i).suffix.lower() == ".pdf" for i in args.inputs) else None
    extractor = f"{exe} ({pdftotext_version(exe)})" if exe else "pre-extracted text"

    docs = []
    for inp in args.inputs:
        p = Path(inp)
        text = pdf_to_text(p, ROOT / "work", exe) if p.suffix.lower() == ".pdf" else p.read_text(encoding="utf-8", errors="replace")
        d = Doc(p.stem, text)
        chars = sum(len(t.strip()) for t in d.pages)
        if not d.pages or chars / max(1, len(d.pages)) < 100:
            sys.exit(f"{p.name}: no usable text layer ({chars} characters over {len(d.pages)} pages). "
                     "This looks like a scanned PDF - run OCR first (e.g. ocrmypdf) and try again.")
        docs.append(d)

    bk = pass_bid_killers(docs)
    hid = pass_hidden(docs)
    ss = pass_sole_source(docs)
    sub = pass_submittals(docs)
    cost = pass_costly(docs, trades, watch)
    rfi = pass_rfi(docs)
    summary = build_summary(docs, bk, hid, ss, sub, cost, rfi, trades)

    method = [{"Item": k, "Detail": v} for k, v in [
        ("What this is", "Auto-generated spec-review form. Every row is a real sentence from the document, cited to Section + PDF page. Verify before relying on it."),
        ("Region", "UK - tender / JCT / NEC wording, NBS & Uniclass sections, GBP, day-month dates. Values are read from the matched sentence only; a value taken from the line after a heading is marked MEDIUM with a note."
                   if args.region == "uk" else "US - CSI MasterFormat and US public-works bid documents."),
        ("PDF page", "Page numbers are PDF page numbers (1 = first sheet of the PDF), not the printed page number in the footer."),
        ("Documents", "; ".join(f"{d.name} ({len(d.pages)} pages, {d.empty_pages} blank/near-blank)" for d in docs)),
        ("Text extractor", extractor),
        ("Confidence", "HIGH = literal clause match (still verify the number/date). MEDIUM = keyword or judgment call - read the clause."),
        ("NOT FOUND", "A bid-killer marked NOT FOUND was not detected. Confirm it truly does not apply - it may be in a document you did not include."),
        ("Hidden Costs", "Clauses that make you carry cost without a pay item: no separate payment, incidental, at Contractor's expense, no cost to Owner, retesting."),
        ("Sole-Source", "GLOBAL clause overrides individual rows. LIMITED LIST = only listed manufacturers without a substitution request. NAMED (confirm) = a brand with no substitution language either way."),
        ("Submittals", "Rows from the spec's own submittal schedule come first, then every sentence requiring a submittal. Claims, notices, RFIs and field sampling are excluded."),
        ("Costly Items", "Trade watchlists in config/watchlists.json. Price drivers only; boilerplate excluded. One row per item with every page it appears on."),
        ("RFI List", "Only ambiguity that moves price: unforeseen conditions, assumed risk, scope set later, subjective/undefined acceptance, match existing, undefined standards. Draft RFI ready to paste."),
        ("Limits", "Text-layer PDFs only (scanned specs need OCR). Tables and drawings are not read. Extraction is high-recall, not zero-noise."),
    ]]

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    write_xlsx(out, [
        ("Summary", summary, "summary"),
        ("Bid-Killers", bk, "table"),
        ("Hidden Costs", hid, "table"),
        ("Sole-Source", ss, "table"),
        ("Submittals", sub, "table"),
        ("Costly Items", cost, "table"),
        ("RFI List", rfi, "table"),
        ("Method & Notes", method, "table"),
    ], multi_doc=len(docs) > 1)

    n = lambda rows, col, none: sum(1 for r in rows if r.get(col) != none)
    print(f"Wrote {out}")
    print(f"  Bid-killers: {sum(1 for r in bk if r['Clause'] != 'NOT FOUND')} of {len(bk)} found")
    print(f"  Hidden costs: {n(hid, 'Hidden cost', 'NONE DETECTED')}")
    print(f"  Named products: {n(ss, 'Classification', 'NONE DETECTED')}")
    print(f"  Submittals: {n(sub, 'Type', 'NONE DETECTED')}")
    print(f"  Costly items: {n(cost, 'Costly item', 'NONE DETECTED')}")
    print(f"  RFIs: {n(rfi, 'Risk', 'NONE DETECTED')}")

if __name__ == "__main__":
    main()
