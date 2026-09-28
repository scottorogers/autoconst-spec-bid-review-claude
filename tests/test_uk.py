#!/usr/bin/env python3
"""
UK-mode checks against a SYNTHETIC tender (tests/fixtures/uk_tender_synthetic.txt) - no PDFs or
pdftotext needed, so this runs anywhere:

    py tests/test_uk.py          (or: python -m pytest tests/test_uk.py)

Every expected value below is written in the fixture file itself - open it to check.
"""
import subprocess, sys, tempfile
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
ENGINE = ROOT / "scripts" / "bid_review.py"
FIXTURE = ROOT / "tests" / "fixtures" / "uk_tender_synthetic.txt"

_wb = None

def wb():
    global _wb
    if _wb is None:
        out = Path(tempfile.mkdtemp()) / "uk.xlsx"
        r = subprocess.run([sys.executable, str(ENGINE), str(FIXTURE), "--out", str(out)],
                           capture_output=True, text=True, encoding="utf-8", errors="replace")
        assert r.returncode == 0, r.stderr
        _wb = load_workbook(out)
    return _wb

def rows(tab):
    data = list(wb()[tab].iter_rows(values_only=True))
    return [dict(zip(data[0], r)) for r in data[1:]]

def bk(label):
    return next(r for r in rows("Bid-Killers") if r["Bid-killer"] == label)

def test_page_column_is_labelled_pdf_page():
    assert "PDF page" in rows("Bid-Killers")[0]

def test_tender_return_date_and_time():
    r = bk("Tender return date / time")
    assert "13 November 2026" in r["Key value"] and "12:00 noon" in r["Key value"] and r["PDF page"] == 1

def test_site_visit_does_not_borrow_next_sentence_value():
    # the next sentence is the 10% performance bond - it must not be reported against the visit
    r = bk("Mandatory site visit / briefing")
    assert "3 November 2026" in r["Key value"] and "%" not in r["Key value"]

def test_queries_deadline():
    assert bk("Tender queries / clarifications deadline")["Key value"] == "6 November 2026"

def test_validity():
    assert bk("Tender validity period")["Key value"] == "90 days"

def test_form_of_contract():
    assert bk("Form of contract / pricing basis")["Key value"] == "JCT Design and Build Sub-Contract 2016"

def test_amendments_found():
    assert "Schedule of Amendments" in bk("Contract amendments")["Clause"]

def test_lads_gbp_per_week():
    r = bk("LADs / delay damages")
    assert r["Key value"] == "£5,000 per week or part thereof" and "per-week" in r["Notes"]

def test_retention_and_bond():
    assert bk("Retention")["Key value"] == "3%; 1.5%"
    assert bk("Performance bond")["Key value"] == "10%"

def test_contract_period_not_taken_from_retention_sentence():
    r = bk("Contract period / completion")
    assert r["Key value"] == "26 weeks" and "Retention" not in r["Clause"]

def test_cdp_picks_the_design_clause():
    assert "Designed Portion" in bk("Design responsibility (CDP)")["Clause"]

def test_accreditations():
    assert bk("Accreditations / competence cards")["Key value"] == "CHAS; Constructionline Gold; NICEIC"

def test_not_in_fixture_is_not_found():
    for label in ("Payment terms", "CDM duties"):
        assert bk(label)["Clause"] == "NOT FOUND"

def test_nbs_section_cited_not_carried_over():
    # page 3 is Y63 - the old engine cited it as the previous page's "§1.1 General"
    sub = rows("Submittals")
    assert sub[0]["Section"].startswith("Y63") and sub[0]["PDF page"] == 3

def test_hidden_costs_uk_wording():
    cats = {r["Hidden cost"] for r in rows("Hidden Costs")}
    assert {"Included in Contract Sum / rates", "No cost to Employer / Client", "Allow for", "Making good"} <= cats

def test_sole_source_uk():
    ss = rows("Sole-Source")
    assert any(r["Product"] == "Unistrut" and r["Classification"].startswith("OR-EQUAL") for r in ss)
    assert any(r["Manufacturer(s)"] == "Unitrunk Ltd" and r["Classification"].startswith("SOLE") for r in ss)
    assert not any(str(r["Product"]).startswith("BS") for r in ss)

def test_rfi_uses_pdf_page_reference():
    rfi = rows("RFI List")
    assert any("as directed by the Contract Administrator" in r["Ambiguous clause"] for r in rfi)
    assert all("PDF p." in r["Draft RFI"] for r in rfi)

def test_uk_watchlists():
    items = {r["Costly item"] for r in rows("Costly Items")}
    assert {"Cable ladder", "Out-of-hours working", "Live / operational environment"} <= items

if __name__ == "__main__":
    tests = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_")]
    failed = 0
    for name, fn in tests:
        try:
            fn()
            print(f"PASS {name}")
        except Exception as e:
            failed += 1
            print(f"FAIL {name} {type(e).__name__} {e}")
    print(f"\n{len(tests) - failed} passed, {failed} failed")
    sys.exit(1 if failed else 0)
