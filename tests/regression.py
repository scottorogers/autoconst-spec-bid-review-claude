#!/usr/bin/env python3
"""
Regression check: re-run bid_review.py on your own real specs and assert hand-checked answers.

    python tests/regression.py            # every spec in SPECS found on disk
    python tests/regression.py myjob      # only these

Spec PDFs are looked up in SPEC_DIR (default: the folder above this repo). Missing PDFs are
SKIPPED; if nothing ran at all the exit code is non-zero. Every pdftotext build on PATH is tried
(xpdf and Poppler emit different text). For checks that need no PDFs run tests/test_uk.py.
"""
import os, subprocess, sys, tempfile
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SPEC_DIR = Path(os.environ.get("SPEC_DIR", ROOT.parent))
ENGINE = ROOT / "scripts" / "bid_review.py"

def rows(wb, tab):
    data = list(wb[tab].iter_rows(values_only=True))
    hdr = data[0]
    return [dict(zip(hdr, r)) for r in data[1:]]

def bk(wb):
    return {r["Bid-killer"]: r for r in rows(wb, "Bid-Killers")}

def found(wb):
    return [r for r in rows(wb, "Bid-Killers") if r["Clause"] != "NOT FOUND"]

# Add your own hand-checked specs here. Keep the PDFs OUT of the repo (put them in SPEC_DIR) and
# keep client names / quoted clauses out of committed files if the repo is shared.
#
# def check_myjob(wb):
#     b = bk(wb)
#     yield "LADs GBP 5,000/week p12", b["LADs / delay damages"]["Key value"].startswith("£5,000") and b["LADs / delay damages"]["PDF page"] == 12
#     yield "retention NOT FOUND", b["Retention"]["Clause"] == "NOT FOUND"
#
# SPECS = {"myjob": ("My Job ITT.pdf", "electrical", check_myjob)}

SPECS = {}

def extractors():
    """Every distinct pdftotext on this machine (xpdf and Poppler emit different text - test both).
    Set PDFTOTEXT to test just one."""
    if os.environ.get("PDFTOTEXT"):
        return [os.environ["PDFTOTEXT"]]
    found, seen = [], set()
    for d in os.environ.get("PATH", "").split(os.pathsep):
        for exe in ("pdftotext.exe", "pdftotext"):
            p = Path(d) / exe
            if p.is_file():
                key = str(p.resolve()).lower()
                if key not in seen:
                    seen.add(key)
                    found.append(str(p))
    return found

def version(exe):
    r = subprocess.run([exe, "-v"], capture_output=True, text=True)
    return ((r.stderr or r.stdout).strip().splitlines() or ["?"])[0]

def main():
    wanted = sys.argv[1:] or list(SPECS)
    exes = extractors()
    if not exes:
        sys.exit("No pdftotext found on PATH.")
    passed = failed = skipped = 0
    with tempfile.TemporaryDirectory() as tmp:
        for exe in exes:
            tag = version(exe)
            print(f"\n=== extractor: {exe} ({tag})")
            env = dict(os.environ, PDFTOTEXT=exe, PYTHONIOENCODING="utf-8")
            for name in wanted:
                pdf, trade, check = SPECS[name]
                src = SPEC_DIR / pdf
                if not src.exists():
                    print(f"SKIP {name}: {src} not found")
                    skipped += 1
                    continue
                out = Path(tmp) / f"{name}-{abs(hash(exe)) % 10**6}.xlsx"
                r = subprocess.run([sys.executable, str(ENGINE), str(src), "--trade", trade, "--out", str(out)],
                                   capture_output=True, text=True, encoding="utf-8", errors="replace", env=env)
                if r.returncode != 0:
                    print(f"FAIL {name}: engine crashed\n{r.stderr[-800:]}")
                    failed += 1
                    continue
                wb = load_workbook(out)
                for label, ok in check(wb):
                    if not ok:
                        print(f"FAIL {name}: {label}")
                    passed += ok
                    failed += (not ok)
                print(f"  {name}: done")
    print(f"\n{passed} passed, {failed} failed, {skipped} spec runs skipped, {len(exes)} extractor(s)")
    if passed + failed == 0:
        # every spec was skipped - that is not a pass
        sys.exit("NOTHING RAN: no spec PDFs found in SPEC_DIR. For a check that needs no PDFs run tests/test_uk.py.")
    sys.exit(1 if failed else 0)

if __name__ == "__main__":
    main()
